from pathlib import Path
from icp_manual_updates import add_updates
from icp_manual_rules import append_rules
from icp_manual_screenshots import add_fresh_screenshot_guide
from icp_manual_examples import append_examples
from icp_manual_navigation import finalize_navigation
from icp_manual_operations import quick_reference, append_operations
from icp_manual_wayfinding import scenario_guide, annotated_details, append_lookup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"C:\Disk F\Aga\EY\TEL\ICP\Projects")
OUT = ROOT / "outputs" / "manual"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "ICP系統操作手冊-v1.5.docx"
SHIP = Path(r"C:\Users\User\AppData\Local\Temp\codex-clipboard-73676402-59d7-4f96-9ec8-fb91495920d9.png")
SCOPE = Path(r"C:\Users\User\AppData\Local\Temp\codex-clipboard-4b7f1bf3-9b33-4f1e-bf72-abecdb662680.png")
SCREENSHOTS = OUT / "screenshots"
ANNOTATED = OUT / "screenshots-annotated"
ANNOTATED.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

styles = doc.styles
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    s = styles[name]
    s.font.name = "Microsoft JhengHei"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    s.font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.15
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)
styles["Heading 1"].font.size = Pt(18)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].font.size = Pt(14)
styles["Heading 2"].font.bold = True
styles["Heading 3"].font.size = Pt(11.5)
styles["Heading 3"].font.bold = True

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table):
    tblPr = table._tbl.tblPr
    tb = tblPr.find(qn("w:tblBorders"))
    if tb is None:
        tb = OxmlElement("w:tblBorders")
        tblPr.append(tb)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        tb.append(el)

def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, "1F4E78")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255,255,255)
            r.font.bold = True
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                shade(cells[i], "F3F6F9")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    # Repeat column labels and keep compact operation cards together.
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = (ri == 0 or (len(rows) <= 4 and ri < len(t.rows)-1))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def step(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

def picture(path, caption, max_height=None):
    if not path.exists():
        return
    picture_width = 7.0
    if max_height is not None:
        with Image.open(path) as source_image:
            image_width, image_height = source_image.size
        picture_width = min(picture_width, max_height * image_width / image_height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(picture_width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9)

def annotation_font(size, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\msjhbd.ttc" if bold else r"C:\Windows\Fonts\msjh.ttc"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()

def annotate_screenshot(source, output, markers, redact_identity=False, crop_height=None):
    image = Image.open(source).convert("RGBA")
    draw = ImageDraw.Draw(image, "RGBA")
    width, height = image.size
    if redact_identity:
        draw.rounded_rectangle((8, height - 122, min(220, width - 8), height - 38), radius=8, fill=(31, 36, 42, 245))
        draw.text((20, height - 92), "Signed-in user", font=annotation_font(16), fill=(230, 233, 236, 255))
        draw.text((20, height - 66), "Identity hidden in guide", font=annotation_font(13), fill=(180, 186, 192, 255))
    for number, x, y, target_x, target_y in markers:
        radius = max(18, int(width * 0.018))
        draw.line((x, y, target_x, target_y), fill=(220, 45, 55, 255), width=max(4, int(width * 0.004)))
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill=(220, 45, 55, 245), outline=(255,255,255,255), width=3)
        label = str(number)
        box = draw.textbbox((0, 0), label, font=annotation_font(radius, bold=True))
        draw.text((x-(box[2]-box[0])/2, y-(box[3]-box[1])/2-2), label, font=annotation_font(radius, bold=True), fill=(255,255,255,255))
    if crop_height is not None and crop_height < height:
        image = image.crop((0, 0, width, crop_height))
    image.convert("RGB").save(output, quality=92)
    return output

def active_menu_center_y(source):
    """Locate the blue highlighted menu row in the left navigation."""
    image = Image.open(source).convert("RGB")
    width, height = image.size
    pixels = image.load()
    scan_width = min(int(width * 0.20), 250)
    scores = []
    for y in range(70, height):
        score = sum(
            1 for x in range(scan_width)
            if 15 <= pixels[x, y][0] <= 45
            and 40 <= pixels[x, y][1] <= 75
            and 65 <= pixels[x, y][2] <= 115
        )
        scores.append((y, score))
    peak_y, peak_score = max(scores, key=lambda item: item[1])
    threshold = max(20, int(peak_score * 0.45))
    matching_rows = {y for y, score in scores if score >= threshold}
    top = peak_y
    bottom = peak_y
    while top - 1 in matching_rows:
        top -= 1
    while bottom + 1 in matching_rows:
        bottom += 1
    return (top + bottom) // 2

title = doc.add_paragraph(style="Title")
title.add_run("ICP 系統操作手冊")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
sub = doc.add_paragraph("功能操作與權限管理指南", style="Subtitle")
sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.add_paragraph("版本 1.5　更新日期 2026 年 10 月 26 日")
doc.add_paragraph(
    "本手冊說明 ICP 系統的導覽、查詢、維護、上傳、報表與權限管理方式。"
    "畫面及可用按鈕會依登入者的角色、資源權限與 DataScope 資料範圍而不同。"
)
doc.add_paragraph("本版於 2026 年 10 月 26 日補拍已登入系統的新版畫面，並依程式與欄位設定整理操作、卡控及成功／失敗教學案例。案例不代表本次已寫入資料；畫面依權限與環境而不同。目錄、章節及卡控入口皆可點擊跳轉。")
doc.add_paragraph("第一次使用先看 [[1.0|工作情境導覽]]；已遇到問題可直達 [[17|錯誤快速索引]]，需要申請異動請看 [[18|管理者申請總表]]。")

doc.add_heading("目錄", level=1)
for item in [
    "1 Quick Start", "2 Common Operations", "3 Functions", "4 Settings", "5 Reports and Data Exchange",
    "6 Languages", "7 Administration", "8 Role Permissions and DataScope", "9 常見問題", "10 Menu Page Index",
    "11 新版畫面操作對照", "12 卡控條件與處理邏輯", "13 上傳案例與錯誤修正", "14 填寫建議與系統卡控的差別",
    "15 完成判斷與異常處理", "16 文件版本與適用範圍", "17 錯誤快速索引", "18 哪些變更需通知管理者"
]:
    toc_entry = doc.add_paragraph(item)
    toc_entry.paragraph_format.space_after = Pt(1)
    toc_entry.paragraph_format.line_spacing = 1.0

doc.add_page_break()
doc.add_heading("1 Quick Start", level=1)
scenario_guide(doc, add_table)
doc.add_page_break()
doc.add_heading("1.1 登入與首頁", level=2)
step("開啟 ICP 網址並完成公司帳號驗證。")
step("登入後確認左下角 Logged in as 顯示正確人員。")
step("從左側選單進入功能；最近使用會列出近期開啟的頁面。")
bullet("若頁面或按鈕未顯示，先確認角色與資源權限。")
bullet("管理者修改權限後，使用者必須重新載入權限或重新登入。")

doc.add_heading("1.2 左側導覽", level=2)
add_table(["區域", "用途"], [
    ("Search menu", "輸入英文頁面名稱快速定位功能。"),
    ("Recent", "快速返回近期使用頁面；可清除紀錄。"),
    ("Pinned", "將常用頁面固定在選單中；按頁面右側的圖釘可釘選，再按一次可取消釘選。"),
    ("Links", "開啟管理者設定的外部或內部連結。"),
    ("Functions", "Ship Info、Add DI/SA、Mass Update 及 Export。"),
    ("Settings", "維護 BU、Warehouse、Delivery To 與各項代碼。"),
    ("Reports", "Shipping Report、Compare ICP vs AR/UR 及 Mass Data Report。"),
    ("Forwarder / Broker", "Forwarder Data Upload、Customs Data Download 及 Customs Clearance Data。"),
    ("Administration", "Users、Resources、Roles 及角色指派。"),
], [1.5, 5.5])
step("在左側選單找到常用頁面，按頁面右側的 Pin 圖示加入 Pinned。")
step("從 Pinned 可直接開啟已固定頁面；若不再需要，按 Unpin／圖釘圖示取消固定。")
bullet("Pinned 只調整個人導覽捷徑，不會新增頁面權限，也不會改變 DataScope。")

nav_source = OUT / "screenshots-20261026" / "40-navigation.jpg"
nav_crop = OUT / "screenshots-20261026-crops" / "navigation.jpg"
nav_crop.parent.mkdir(exist_ok=True)
with Image.open(nav_source) as nav_image:
    nav_image.crop((0, 0, 900, min(800, nav_image.height))).save(nav_crop)
picture(nav_crop, "實拍局部 左側釘選與最近使用 2026-10-26", max_height=3.4)
doc.add_paragraph("Search menu 搜尋選單；Pinned／釘選保留常用頁，Recent／最近使用列出近期開啟頁面。按圖釘只調整導覽，不會變更權限。")
doc.add_heading("1.3 Language", level=2)
language_source = OUT / "screenshots-20261026" / "41-language-menu.jpg"
language_crop = nav_crop.with_name("language.jpg")
with Image.open(language_source) as language_image:
    language_image.crop((max(0,language_image.width-650),0,language_image.width,350)).save(language_crop)
picture(language_crop, "實拍局部 語言選單 2026-10-26", max_height=2.0)
doc.add_paragraph("右上角語言可選繁體中文、English、日本語；切換不改變資料及權限。本版截圖以繁體中文介面為主，保留部分英文按鈕名稱供對照。")

doc.add_page_break()
quick_reference(doc, add_table)
doc.add_page_break()
doc.add_heading("2 Common Operations", level=1)
doc.add_heading("2.1 清單查詢與篩選", level=2)
step("進入清單頁後，先確認頁面標題與資料筆數。")
step("按欄位標題可切換排序；按欄位下方的「篩選」開啟條件。")
step("選擇或輸入條件後套用；按目前篩選數量可再次調整。")
step("使用右下角頁次及每頁筆數切換資料。")
bullet("畫面查無資料時，可能是篩選條件、DataScope 或角色權限造成。")
bullet("匯出內容應與目前權限及資料範圍一致。")

doc.add_heading("2.2 新增、編輯與刪除", level=2)
bullet("新增：按新增按鈕，填寫必填欄位後儲存。")
bullet("編輯：在清單的操作欄按編輯，確認資料後儲存。")
bullet("刪除：僅在確認資料不再使用時執行；若系統要求再次確認，核對目標後再提交。")
bullet("上傳：先核對檔案格式、工作表名稱與欄位，再執行上傳。")

add_updates("common", doc, add_table, step, bullet, ROOT)

doc.add_heading("3 Functions", level=1)
doc.add_heading("3.1 Ship Info", level=2)
doc.add_paragraph("Ship Info 頁分為 Ship Header 與 Ship Detail。先選取 Header 列，系統才會載入對應 Detail，並啟用 Deposit 或 ARUR Case 等可用操作。")

step("在 Ship Header 欄位使用 Filter 或排序找到目標 Invoice No。")
step("點選資料列左側單選按鈕，載入該筆 Ship Detail。")
step("需要修改時，按 Action 欄的 Edit 圖示。")
step("核對表頭、明細與附件後儲存。")
bullet("若直接網址或 API 指向無權限的 Header，系統仍應拒絕存取。")
bullet("BROKER、BU、FORWARDER、MDP FLAG 等欄位可能受 DataScope 限制。")

add_updates("ship", doc, add_table, step, bullet, ROOT)

for heading, purpose, steps in [
    ("3.2 Add DI/SA", "建立新的 DI/SA 資料。", ["進入 Add DI/SA。", "下載新版 Sample 並填寫 Upload Data。", "上傳驗證、核對完整預覽，再按 Save。"]),
    ("3.3 Mass Update Non NCPI", "以批次檔案更新非 NCPI 資料。", ["下載或準備規定格式。", "選擇檔案並執行驗證。", "確認預覽與錯誤列後提交更新。"]),
    ("3.4 Mass Update NCPI", "以批次檔案更新 NCPI 資料。", ["準備 NCPI 專用格式。", "上傳並檢查驗證結果。", "確認影響範圍後提交。"]),
    ("3.5 Export", "上傳 Shipping advice Excel，透過 Hub 與 FileGenerator 產生檔案。", ["選擇 Shipping advice Excel。", "核對檔案後按上傳並通知 Hub，此動作會建立背景產檔工作。", "查看工作狀態與錯誤；Completed 後按檢視檔案，逐檔下載或全部下載 ZIP。"]),
]:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(purpose)
    for s in steps: step(s)
    add_updates(heading.split()[0], doc, add_table, step, bullet, ROOT)

doc.add_heading("3.6 通知信獨立頁面", level=2)
doc.add_paragraph("從系統管理或已釘選的通知信開啟。上方 SMTP 與各類通知設定僅供查看；更改啟用、收件者或時段由管理者維護 appsettings，本頁不提供直接編輯。")
step("在下方選取到貨延遲、Arrival notice、管制品到貨通知或押金平台起案分頁。")
step("使用欄位 Filter 查詢 Invoice、時間或狀態；點預覽核對信件及錯誤。")
step("預覽不會寄信；狀態與寄送規則見第 12 章。範例畫面的啟用與收件者只是目前環境，不應照抄到其他主機。")
doc.add_heading("通知信使用者須知與變更申請", level=3)
doc.add_paragraph("需要新增、移除或修改通知信 MailTo（主要收件者）／CcTo（副本收件者）時，請通知系統管理者更新。Notification Mail 頁面的共用設定僅供查看，使用者不能在此直接修改；也不要把預覽畫面當作已完成設定變更。")
add_table(["項目", "使用者應如何處理"], [
    ("共用收件者與副本", "延遲交貨通知、管制品到貨通知及押金起案通知的共用 MailTo／CcTo，由管理者維護。到貨通知的共用 CcTo 亦由管理者維護。請明確說明新增、移除或取代，不要只提供一份無異動說明的名單。"),
    ("到貨通知的個別收件者", "到貨通知實際 To 取自 Ship Info 的 Arrival Notice 欄位；具編輯權限者可依流程修改，多個地址以分號分隔。儲存後重新核對待寄排程與預覽；此欄不會修改共用 CcTo。沒有編輯權限時，請交由有權限人員處理。"),
    ("寄件者與寄送時段", "寄件者、SMTP、通知啟用狀態、寄送頻率與時段均請管理者處理。一般使用者不需自行修改設定檔。沒有指定寄件者時，系統預設使用 TETIS87GEMSAdmin@tel.com；現有有效寄件者設定會保留。"),
    ("申請時提供的資訊", "提供環境（測試／正式及網址）、通知種類、原收件者、新增／移除後的 MailTo 與 CcTo、預定生效時間、申請原因及聯絡人。若僅影響個別到貨通知，另附 Invoice No。"),
    ("變更完成後的核對", "請管理者確認已更新正確環境及設定已生效，再重新開啟設定頁與相關郵件預覽核對 To／Cc。既有待寄排程與已寄紀錄不應假設會全部同步改寫；需要調整特定待寄通知時，請一併提供 Invoice 與排程時間，由管理者確認處理方式。"),
    ("寄送結果與異常", "Pending 為待寄、Sending 為寄送中、Sent 為系統已完成寄送作業、Failed 為失敗。Sent 不保證對方信箱已收件。未收到信時先核對收件者、狀態、排程時間與垃圾郵件，再提供 Invoice、通知種類、發生時間及錯誤畫面給管理者；不要反覆建立通知，以免重複寄送。")
], [1.4, 5.6])
doc.add_paragraph("變更申請範例：正式環境／延遲交貨通知，請自 2026-10-27 起在 MailTo 新增 user@example.com、在 CcTo 移除 old@example.com，其餘收件者與排程保持不變。以上地址僅為教學範例，請填入實際核准地址。")

doc.add_heading("4 Settings", level=1)
doc.add_paragraph("設定頁供管理者維護系統選項。異動會影響其他頁面的下拉選單、預設值或計算邏輯，儲存前應先確認代碼是否已被使用。")
add_table(["頁面", "管理內容", "基本操作"], [
    ("BU Codes", "BU 代碼與說明", "Search, Add, Edit"),
    ("Warehouse Codes", "倉庫代碼", "Search, Add, Edit"),
    ("Delivery To List", "交貨對象資料", "Search, Add, Edit"),
    ("Pick Up Location", "提貨地點資料", "Search, Add, Edit"),
    ("ETA Delivery Date Table", "ETA 與交貨日規則", "Search, Add, Edit"),
    ("Default Delivery Warehouse", "預設倉庫對應", "Search, Add, Edit"),
    ("Order Type", "訂單類型選項", "Search, Add, Edit"),
    ("Air Sea", "Air／Sea 選項", "Search, Add, Edit"),
    ("Broker", "Broker 選項", "Search, Add, Edit"),
    ("Invoice Type", "Invoice Type 選項", "Search, Add, Edit"),
    ("Order Priority", "Order Priority 選項", "Search, Add, Edit"),
    ("Customized", "系統自訂參數", "Search and Maintain"),
    ("Shipper", "發貨方設定", "Search, Add, Edit"),
    ("Country of Origin", "原產國設定", "Search, Add, Edit"),
], [1.6, 3.0, 2.4])

add_updates("settings", doc, add_table, step, bullet, ROOT)

doc.add_heading("5 Reports and Data Exchange", level=1)
for heading, body, notes in [
    ("5.1 Shipping Report", "依條件查詢及匯出船務表頭與相關明細。", "清單、篩選選項與 Excel 匯出均套用 DataScope。"),
    ("5.2 Compare ICP vs AR/UR", "比對 ICP 與 AR/UR 資料差異。", "本頁目前不套用 DataScope。"),
    ("5.3 Mass Data Report", "查詢及匯出大量船務資料。", "清單、篩選與匯出均套用 DataScope。"),
    ("5.4 Forwarder Data Upload", "上傳貨代提供的資料檔案。", "本頁目前不套用 DataScope；上傳前先核對模板。"),
    ("5.5 Customs Data Download", "依條件下載報關所需資料。", "下載內容依登入者的 DataScope 限制。"),
    ("5.6 Customs Clearance Data", "查詢、維護及匯出關稅資料與附件。", "清單、篩選、附件與匯出均套用 DataScope。"),
]:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(body)
    if heading.startswith(("5.2", "5.5")):
        step("設定查詢條件，核對結果後下載。")
    add_updates(heading.split()[0], doc, add_table, step, bullet, ROOT)
    doc.add_paragraph("注意：" + notes)

doc.add_heading("6 Languages", level=1)
doc.add_paragraph("Languages 用於維護繁體中文、英文與日文顯示文字。")
step("進入 Languages 並搜尋資源鍵。")
step("編輯各語言文字，保持相同資源鍵的語意一致。")
step("儲存後重新載入頁面，確認畫面顯示。")
bullet("不要直接更改資源鍵；程式會以資源鍵尋找翻譯。")

doc.add_heading("7 Administration", level=1)
add_table(["頁面", "用途"], [
    ("Users", "查詢使用者並查看 Permissions。"),
    ("Resources", "管理頁面、Action 與 ResourceCode。"),
    ("Roles", "建立及維護角色。"),
    ("Role Permissions", "將 Roles、Resources 與 DataScope 組合。"),
    ("User Roles", "將角色指派給特定 TELID。"),
    ("Department Roles", "依部門指派角色。"),
    ("Mail Group Roles", "依郵件群組指派角色。"),
    ("Notification Mail", "檢視通知設定、各類清單、郵件預覽與錯誤訊息；設定由管理者維護 appsettings。"),
], [1.7, 5.3])
doc.add_heading("7.1 建議設定順序", level=2)
step("先確認資源已存在且 ResourceCode、Route 與 Action 正確。")
step("建立角色。")
step("在 Role Permissions 選取角色與資源，必要時設定 DataScope。")
step("透過人員、部門或郵件群組指派角色。")
step("讓目標使用者重新載入權限或重新登入，然後測試功能。")

doc.add_heading("8 Role Permissions and DataScope", level=1)
doc.add_paragraph("DataScope 用來限制同一功能中可見的資料。未勾選 Restrict visible data 時不限制；勾選後依資料表、欄位、運算子與值建立條件。")
picture(SCOPE, "圖 2　Role Permissions 建立流程中的 Data scope 設定")
doc.add_heading("8.1 設定步驟", level=2)
step("在 Role Permissions 精靈選擇角色。")
step("選擇要授權的資源。")
step("進入 Set data scope，勾選 Restrict visible data。")
step("選擇資料表、欄位及運算子，再輸入值。")
step("需要多個欄位時按 Add condition；同一規則中的條件必須全部符合。")
step("確認摘要並建立權限，最後重新載入使用者權限。")
add_table(["運算子", "Value 範例", "結果"], [
    ("=", "LOG YUANFAN", "只允許完全等於此值。"),
    ("in", "LOG YUANFAN, LOG KWE", "允許符合其中任一值；不需加單引號。"),
    ("contains", "YUANFAN", "允許欄位中包含此文字。"),
], [1.0, 2.5, 3.5])
doc.add_heading("8.2 套用頁面", level=2)
add_table(["頁面", "常用資料表", "常用欄位"], [
    ("Ship Info", "ICP_HEADER", "BROKER、BU、FORWARDER、MDP_FLAG"),
    ("Shipping Report", "ICP_HEADER", "BROKER、BU、FORWARDER"),
    ("Mass Data Report", "ICP_HEADER", "BROKER、BU、FORWARDER"),
    ("Customs Clearance Data", "TariffData", "Broker、AirSea、Shipper"),
    ("Customs Data Download", "實際下載資料來源", "依下載資料欄位設定"),
], [1.8, 2.0, 3.2])
doc.add_paragraph("目前系統規則：DataScope JSON、資料表、欄位、型別或運算子無法解析時，該規則視為未設定，因此不限制資料。管理者應在正式授權前使用目標帳號驗證。")

doc.add_heading("9 常見問題", level=1)
add_table(["現象", "檢查方式"], [
    ("看不到頁面", "確認資源權限、角色指派與 Route；重新登入。"),
    ("修改權限後沒有變化", "重新載入 Session 權限或登出再登入。"),
    ("查不到資料", "清除欄位篩選，確認 DataScope 的資料表、欄位與值。"),
    ("DataScope 顯示全部資料", "檢查 JSON、資料表、欄位、運算子與值是否有效；錯誤規則目前視為未限制。"),
    ("Ship Info 明細未出現", "先選取一筆 Ship Header；確認該 Header 在可見範圍。"),
    ("匯出筆數不同", "確認匯出條件、目前篩選與 DataScope 是否一致。"),
    ("上傳失敗", "檢查檔案格式、必填欄位、重複資料及畫面錯誤訊息。"),
    ("上傳通過但未新增資料", "確認是否已按 Save；預覽成功不代表正式儲存。"),
    ("Upload Data 無法辨識", "確認第一個分頁名稱為 Upload Data、第一列保留欄位標題，第二列開始填資料。"),
    ("Delivery To 驗證失敗", "Add DI/SA 使用 Delivery To List 的 Customer Code，並非 Key 或 Name。"),
    ("歷史紀錄顯示 File not found", "已匯入紀錄仍存在，原檔目前不可取得；請管理者檢查檔案保存位置。"),
    ("Mass Data Report 沒有 Detail", "目前設計僅保留 Header；下載內容與 Header 一致。"),
    ("Settings Excel 上傳失敗", "原資料保留；依錯誤修正所有列後重傳，不要用空白檔案清空資料。"),
], [2.3, 4.7])

doc.add_heading("10 Menu Page Index", level=1)
pages = [
    ("Functions", "Ship Info", "/ShipInfo"), ("Functions", "Add DI/SA", "/AddDiSa"),
    ("Functions", "Mass Update Non NCPI", "/MassUpdateNonNcpi"), ("Functions", "Mass Update NCPI", "/MassUpdateNcpi"),
    ("Functions", "Export", "/Export"), ("Settings", "BU Codes", "/BuCode"), ("Settings", "Warehouse Codes", "/WhCode"),
    ("Settings", "Delivery To List", "/DeliveryToList"), ("Settings", "Pick Up Location", "/PickUpLocation"),
    ("Settings", "ETA Delivery Date Table", "/EtaDelDateTable"), ("Settings", "Default Delivery Warehouse", "/DefaultDeliveryWh"),
    ("Settings", "Order Type", "/OrderType"), ("Settings", "Air Sea", "/AirSea"), ("Settings", "Broker", "/Broker"),
    ("Settings", "Shipper", "/Shipper"), ("Settings", "Country of Origin", "/CountryOfOrigin"),
    ("Settings", "Invoice Type", "/InvoiceType"), ("Settings", "Order Priority", "/OrderPriority"), ("Settings", "Customized", "/Customized"),
    ("Reports", "Shipping Report", "/ShippingReport"), ("Reports", "Compare ICP vs AR/UR", "/CompareIcpVsArUr"),
    ("Reports", "Mass Data Report", "/MassDataReport"), ("Forwarder", "Forwarder Data Upload", "/ForwarderDataUpload"),
    ("Broker", "Customs Data Download", "/CustomsDataDownload"), ("Broker", "Customs Clearance Data", "/TariffData"),
    ("Languages", "Languages", "/LocalizationManagement"), ("Administration", "Users", "/Users"),
    ("Administration", "Resources", "/Resources"), ("Administration", "Roles", "/Roles"), ("Administration", "Role Permissions", "/RolePermissions"),
    ("Administration", "User Roles", "/RoleTelIds"), ("Administration", "Department Roles", "/RoleDepIds"),
    ("Administration", "Mail Group Roles", "/RoleMailGroups"),
    ("Administration", "Notification Mail", "/NotificationMail"),
]
add_table(["Category", "Page", "Route"], pages, [1.5, 3.3, 2.2])

add_fresh_screenshot_guide(doc, add_table, picture, ROOT)
annotated_details(doc, picture, ROOT)

append_rules(doc, add_table, bullet, ROOT)
append_examples(doc, add_table, bullet, ROOT)
append_operations(doc, add_table, bullet)
append_lookup(doc, add_table)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("ICP 系統操作手冊 版本 1.5")

doc.core_properties.title = "ICP 系統操作手冊"
doc.core_properties.subject = "ICP 功能操作與權限管理"
doc.core_properties.author = "ICP Project Team"
finalize_navigation(doc)
doc.save(DOCX)
print(DOCX)
