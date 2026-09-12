from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Disk F\Aga\EY\TEL\ICP\Projects")
OUT = ROOT / "outputs" / "manual"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "ICP系統設定管理手冊.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    style.font.name = "Microsoft JhengHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(10.2)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Normal"].paragraph_format.line_spacing = 1.12
doc.styles["Title"].font.size = Pt(27)
doc.styles["Title"].font.bold = True
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
border = title_ppr.find(qn("w:pBdr"))
if border is not None:
    title_ppr.remove(border)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 2"].font.size = Pt(13.5)
doc.styles["Heading 2"].font.bold = True

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)

def set_borders(table):
    props = table._tbl.tblPr
    borders = props.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[i], "F3F6F9")
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def step(text):
    doc.add_paragraph(text, style="List Number")

def code_block(lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        run.font.size = Pt(8.5)

doc.add_paragraph("ICP 系統設定管理手冊", style="Title")
doc.add_paragraph("完整 appsettings 設定與部署檢查指南", style="Subtitle")
doc.add_paragraph("版本 1.0　更新日期 2026 年 9 月 12 日")
doc.add_paragraph(
    "本手冊提供 ICP 管理與部署人員使用，逐節說明 appsettings.example.json 內的所有設定，"
    "以及資料庫、RabbitMQ、上傳路徑與共用資料夾的部署方式。範例不包含正式帳號、密碼、"
    "伺服器名稱或其他機密值。"
)

doc.add_heading("1 Configuration Files", level=1)
add_table(["File", "Purpose", "Recommendation"], [
    ("appsettings.json", "系統共用預設值。", "只放所有環境都適用的設定；避免保存正式密碼。"),
    ("appsettings.Development.json", "Development 環境覆寫。", "只覆寫開發環境需要不同的欄位。"),
    ("appsettings.TEL.json", "TEL 部署環境設定。", "由部署流程保護，不要提交真實密碼。"),
    ("appsettings.TEL.<Machine>.json", "特定主機覆寫。", "僅放該主機需要不同的值。"),
    ("appsettings.example.json", "可提交的設定範本。", "新增設定欄位時同步更新並使用假值。"),
], [2.55, 2.05, 2.4])
doc.add_paragraph("ASP.NET Core 會依環境與設定來源覆寫同名欄位。實際生效值取決於啟動環境、檔案載入順序、環境變數及部署平台設定。")

doc.add_heading("2 App", level=1)
add_table(["Key", "Description", "Operational note"], [
    ("BrandName", "網站品牌名稱。", "修改後檢查頁首與登入頁。"),
    ("TitleSuffix", "瀏覽器標題後綴。", "修改後重新整理頁面。"),
    ("Mode", "驗證模式，例如 PRD 或 DEV。", "正式環境必須使用核准模式。"),
    ("SimulatedWindowsIdentity", "本機模擬 Windows 身分。", "只可用於受控開發環境。"),
    ("SuperUser", "是否允許切換其他使用者。", "正式環境必須保持 Off，除非經正式核准。"),
    ("DevUser", "DEV 模式的模擬使用者資訊。", "不得使用真實個資作為可提交範例。"),
], [2.1, 3.1, 2.0])

doc.add_heading("3 Sidebar", level=1)
add_table(["Key", "Description", "Allowed or example"], [
    ("RecentLimit", "Recent 顯示筆數。", "0 到 50。"),
    ("RecentEnabled", "是否顯示 Recent。", "true / false。"),
    ("PinnedLimit", "Pinned 可保存筆數。", "0 到 20。"),
    ("DefaultPage", "登入後預設 Controller。", "例如 ShipInfo；必須是允許的 ICP 頁面。"),
    ("MenuOrder", "各選單頁的排序權重。", "較小值排在前面，範圍 -10000 到 10000。"),
    ("CustomLinks", "Links 區塊的自訂連結。", "支援 / 開頭站內路徑或 http / https 網址。"),
], [2.0, 3.0, 2.2])
doc.add_heading("3.1 CustomLinks fields", level=2)
add_table(["Field", "Description"], [
    ("Name", "側邊選單顯示名稱；啟用的名稱不可重複。"),
    ("Url", "站內根相對路徑或完整 http / https URL；啟用的 URL 不可重複。"),
    ("OpenInNewTab", "true 時以新分頁開啟。"),
    ("Enabled", "false 時不顯示連結。"),
    ("SortOrder", "Links 內的排序值，較小值在前。"),
], [2.1, 5.1])
doc.add_paragraph("文件連結範例：Name = ICP Configuration Guide，Url = /docs/ICP-Configuration-Guide.docx。")

doc.add_heading("4 ConnectionStrings", level=1)
add_table(["Key", "Database purpose", "Required"], [
    ("ILC_Connection", "使用者、部門與組織資料。", "是"),
    ("ICP_Connection", "ICP 業務資料、設定與權限資料。", "是"),
    ("FIESTA_Connection", "FIESTA 或郵件群組整合資料。", "使用相關功能時"),
], [2.1, 3.7, 1.4])
doc.add_heading("4.1 Connection string example", level=2)
doc.add_paragraph("下列只示範格式；請由 Secret、環境變數或受保護的主機設定提供真實值。")
code_block([
    '"ConnectionStrings": {',
    '  "ICP_Connection": "Server=<DB_SERVER>;Database=<DB_NAME>;User Id=<SERVICE_ACCOUNT>;Password=<SECRET>;Encrypt=True;TrustServerCertificate=False;MultipleActiveResultSets=True"',
    '}'
])
bullet("不要將正式帳號、密碼或完整連線字串提交到 Git。")
bullet("優先使用部署平台 Secret、環境變數或受保護的主機設定檔。")
bullet("修改後需重新啟動服務，並測試登入、清單查詢及資料庫寫入。")
bullet("正式環境應啟用加密連線並依組織政策設定憑證驗證。")
doc.add_heading("4.2 Database validation and troubleshooting", level=2)
add_table(["Symptom", "Check"], [
    ("Login failed", "帳號、密碼、驗證模式、資料庫使用者 mapping 與最小必要權限。"),
    ("Server not found or timeout", "DNS、Port、Firewall、SQL Server instance 與網路路由。"),
    ("Certificate error", "Encrypt、TrustServerCertificate 與伺服器憑證鏈。正式環境不要用略過驗證作為長期修正。"),
    ("Only one function fails", "確認該功能使用的是 ILC、ICP 或 FIESTA 哪一條連線。"),
], [2.2, 5.0])

doc.add_heading("5 Upload Paths and Shared Folders", level=1)
add_table(["Section", "Key", "Description"], [
    ("ForwarderDataUpload", "StoragePath", "貨代上傳檔案的保存位置。"),
    ("ForwarderDataUpload", "MaxSizeMb", "單一上傳檔案大小上限。"),
    ("TariffData", "StoragePath", "Tariff Data 檔案保存位置。"),
    ("TariffData", "MaxSizeMb", "Tariff Data 上傳大小上限。"),
    ("TariffData", "BrokerKeywords", "檔名或內容辨識 Broker 的關鍵字對應。"),
    ("ArurAttachment", "IpcStorageRoot", "ARUR 附件根目錄。"),
    ("ArurAttachment", "RelativeFolder", "根目錄下的相對資料夾。"),
    ("ArurAttachment", "MaxSizeMb", "附件大小上限。"),
    ("ArurAttachment", "AllowedExtensions", "允許的副檔名白名單。"),
    ("Integration.Export", "OutputDirectory", "ICP 與 FileGenerator 共用的輸出根目錄。"),
], [2.1, 2.0, 3.1])
doc.add_heading("5.1 Path examples", level=2)
code_block([
    '"ForwarderDataUpload": { "StoragePath": "D:\\\\ICPData\\\\Forwarder", "MaxSizeMb": 50 },',
    '"TariffData": { "StoragePath": "\\\\\\\\FILESERVER\\\\ICP\\\\Tariff", "MaxSizeMb": 50 },',
    '"ArurAttachment": {',
    '  "IpcStorageRoot": "\\\\\\\\FILESERVER\\\\ICP",',
    '  "RelativeFolder": "ATTACHED_FILE",',
    '  "MaxSizeMb": 50,',
    '  "AllowedExtensions": [".pdf", ".xlsx", ".zip"]',
    '},',
    '"Integration": { "Export": { "OutputDirectory": "\\\\\\\\FILESERVER\\\\ICP\\\\Export" } }'
])
doc.add_paragraph("相對路徑會以應用程式 ContentRoot 為基準；正式環境或多台主機共用檔案時，建議使用 UNC 路徑。Windows Service 或 IIS App Pool 不應依賴登入使用者的 mapped drive，例如 Z:。")
doc.add_heading("5.2 Folder permissions", level=2)
bullet("以實際執行 ICP 的 Windows Service 或 IIS Application Pool 身分授權。")
bullet("UNC 目錄必須同時具有 Share permission 與 NTFS permission；一般上傳與輸出需要 Read、Write、Modify。")
bullet("只授予所需資料夾與權限，不要使用 Everyone Full Control。")
bullet("IpcStorageRoot 與 RelativeFolder 組合後會再依 Invoice No 建立子目錄；根目錄必須存在且可存取。")
doc.add_heading("5.3 Upload path validation", level=2)
step("使用應用程式實際服務身分建立、讀取並刪除測試檔案。")
step("測試 Forwarder Data Upload、Tariff Data 上傳，以及 Ship Info 附件上傳與下載。")
step("確認 Integration.Export.OutputDirectory 與 FileGenerator 的 OutputDirectory 指向同一個實體目錄。")
step("檢查 Access denied、Path not found、檔案鎖定、防毒隔離與磁碟空間錯誤。")
step("調整 MaxSizeMb 時，同步檢查 IIS、反向代理與應用程式的 Request Size 限制。")

doc.add_heading("6 Localization and Table Fields", level=1)
add_table(["Section", "Description"], [
    ("LocalizationManagement.Enabled", "控制 Languages 管理功能是否啟用。"),
    ("ShipInfoProDataTableFields", "Ship Info 及相關報表顯示欄位設定。"),
    ("ForwarderTableFields", "Forwarder Data Upload 資料表欄位設定。"),
    ("CustomsDataDownloadTableFields", "Customs Data Download 顯示或下載欄位設定。"),
    ("TariffTableFields", "Tariff Data 資料表欄位設定。"),
], [3.0, 4.2])
doc.add_paragraph("欄位設定通常會影響顯示順序、名稱、是否可見、篩選或匯出。修改後要逐頁檢查表頭、Filter、分頁與 Excel 輸出。")
doc.add_paragraph("appsettings.example.json 的 _comment_ShipInfoProDataTableFields 指向 Config/shipinfo-pro-datatable-fields.json；這是說明欄位，不會直接成為執行設定。_comment_IntegrationEvents 同樣只提供文件位置。")

doc.add_heading("7 Integration", level=1)
add_table(["Section", "Purpose", "Key checks"], [
    ("RabbitMq", "發佈 ICP 整合事件。", "Enabled、HostName、Port、VirtualHost、Exchange、RoutingKey 與認證。"),
    ("Outbox", "控制待發事件的輪詢與重試。", "PollIntervalSeconds、MaxRetryCount、BatchSize。"),
    ("Hub", "TEL Integration Hub 位置。", "BaseUrl 可連線且使用正確環境。"),
    ("Export", "FileGenerator 共用輸出位置。", "OutputDirectory 存在且服務帳號可讀寫。"),
], [1.7, 2.3, 3.2])
bullet("RabbitMQ 密碼屬於 Secret，不應保存在可提交檔案。")
bullet("啟用整合前先確認目的環境、交換器與 Routing Key，避免將測試事件送往正式系統。")
doc.add_heading("7.1 RabbitMQ settings", level=2)
add_table(["Key", "Description"], [
    ("Enabled", "false 時不發佈 RabbitMQ 訊息；確認連線與目的環境後才設為 true。"),
    ("HostName / Port", "Broker 主機與 AMQP Port，常見未加密連線為 5672；仍以環境規格為準。"),
    ("VirtualHost", "RabbitMQ virtual host；使用者必須被授權存取。"),
    ("UserName / Password", "服務帳號與密碼；Password 必須由 Secret 或受保護的環境設定供應。"),
    ("Exchange", "事件發佈的 exchange；必須與 RabbitMQ 端拓樸一致。"),
    ("RoutingKey", "事件 routing key；目前範例為 icp.shipinfo.case.initiated。"),
], [2.2, 5.0])
code_block([
    '"RabbitMq": {',
    '  "Enabled": true, "HostName": "<RABBITMQ_HOST>", "Port": 5672,',
    '  "VirtualHost": "<VHOST>", "UserName": "<SERVICE_ACCOUNT>",',
    '  "Password": "<SECRET>", "Exchange": "tel.integration",',
    '  "RoutingKey": "icp.shipinfo.case.initiated"',
    '}'
])
doc.add_heading("7.2 Outbox Hub and Export", level=2)
add_table(["Key", "Description"], [
    ("Outbox.PollIntervalSeconds", "背景服務掃描待發事件的秒數間隔。"),
    ("Outbox.MaxRetryCount", "單筆事件最大重試次數。"),
    ("Outbox.BatchSize", "每批處理筆數；調高會增加瞬間負載。"),
    ("Hub.BaseUrl", "TEL Integration Hub 的基底 URL，必須指向正確環境。"),
    ("Export.OutputDirectory", "與 FileGenerator 共用的輸出位置；路徑與權限見第 5 章。"),
], [2.7, 4.5])
doc.add_paragraph("驗證 RabbitMQ 時，先確認 DNS 與 Port，再確認帳密、VirtualHost permission、Exchange 與 Routing Key。啟用後建立一筆測試事件，檢查 Outbox 是否送出，並查看應用程式 Log 是否出現 connection refused、authentication refused 或 exchange not found。")

doc.add_heading("8 Logging and Hosting", level=1)
add_table(["Section", "Description", "Recommendation"], [
    ("Logging.LogLevel", "ASP.NET Core 記錄層級。", "正式環境避免長期使用 Debug 或 Trace。"),
    ("Serilog.MinimumLevel", "Serilog 全域與來源層級。", "保留必要診斷資訊，避免記錄敏感內容。"),
    ("Serilog.WriteTo", "Log 輸出目的地與格式。", "確認路徑、保留天數、檔案大小及磁碟空間。"),
    ("AllowedHosts", "允許的 Host header。", "正式環境設定明確主機名稱。"),
], [2.0, 2.7, 2.5])
doc.add_paragraph("AllowedHosts = * 代表不限制 Host header；正式環境建議改成核准的主機名稱。LogLevel 與 Serilog 設定若同時存在，需確認實際 provider 與 override 層級，不要在正式環境長期記錄 Debug、Trace、密碼或完整連線字串。")

doc.add_heading("9 Change Procedure", level=1)
for number, instruction in enumerate([
    "先確認要修改的環境與實際載入的 appsettings 檔案。",
    "備份原設定，但不要把含 Secret 的備份放入網站目錄或 Git。",
    "只修改必要欄位，保持 JSON 結構、逗號與型別正確。",
    "重新啟動服務；少數由 IOptionsMonitor 或 IConfiguration 讀取的設定可熱更新，但正式變更仍建議重啟驗證。",
    "依變更範圍執行登入、Menu、查詢、上傳、匯出與整合測試。",
    "檢查應用程式 Log，確認沒有 Options validation、資料庫、路徑或連線錯誤。",
], start=1):
    doc.add_paragraph(f"{number}.　{instruction}")

doc.add_heading("10 Deployment Checklist", level=1)
for item in [
    "App.Mode 與 SuperUser 符合正式環境政策。",
    "ConnectionStrings 指向正確環境且 Secret 未提交 Git。",
    "StoragePath、OutputDirectory 與附件目錄存在並具有適當權限。",
    "CustomLinks 使用可部署的站內路徑或核准的 https URL。",
    "MenuOrder、DefaultPage 與角色資源權限一致。",
    "RabbitMq、Hub 與 Outbox 指向正確環境並完成連線測試。",
    "Logging 層級、檔案保留與磁碟容量已確認。",
    "重新登入並完成關鍵頁面的 Smoke Test。",
]:
    bullet(item)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("ICP Configuration Guide 1.0")
doc.core_properties.title = "ICP 系統設定管理手冊"
doc.core_properties.subject = "appsettings 設定與部署檢查"
doc.core_properties.author = "ICP Project Team"
doc.save(DOCX)
print(DOCX)
