"""Internal document links and page fields for the ICP user guide."""
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER


def link(paragraph, text, anchor):
    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor)
    h.set(qn('w:history'), '1')
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    r.append(props); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r)
    paragraph._p.append(h)


def field(paragraph, instruction):
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instruction)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '1'; r.append(t); fld.append(r)
    paragraph._p.append(fld)


def bookmark(paragraph, anchor, number):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(number)); start.set(qn('w:name'), anchor)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(number))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def finalize_navigation(doc):
    headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
    targets = {}
    for i, p in enumerate(headings, 1):
        anchor = 'icp_section_' + str(i)
        bookmark(p, anchor, i)
        targets[p.text] = anchor
    contents = next(p for p in headings if p.text == '目錄')
    contents_anchor = targets['目錄']
    rules_anchor = next(a for t,a in targets.items() if t.startswith('12 卡控'))
    numbered = {m.group(1): a for t,a in targets.items() if (m := re.match(r'^(\d+(?:\.\d+)?)\s',t))}
    rules_for = {'3.1':'12.5','3.2':'12.1','3.6':'12.7','4':'12.4','5.1':'12.6','5.3':'12.6','5.4':'12.2','5.6':'12.3'}
    cases_for = {'3.2':'13.1','3.3':'13.2','3.4':'13.3','3.5':'13.12','4':'13.9','5.4':'13.4','5.6':'13.5','6':'13.11'}
    case_rules = {'13.1':'12.1','13.2':'15.2','13.3':'15.2','13.4':'12.2','13.5':'12.3',
                  '13.6':'5.6','13.7':'5.6','13.8':'5.6','13.9':'12.4','13.10':'12.5','13.11':'6','13.12':'3.5'}
    rules_for.update(case_rules)
    cases_for.update({'12.1':'13.1','12.2':'13.4','12.3':'13.5','12.4':'13.9','12.5':'13.10','15.2':'13.2'})
    # Replace the existing contents entries with internal hyperlinks and page refs.
    in_contents = False
    for p in doc.paragraphs:
        if p._p is contents._p:
            in_contents = True
            continue
        if in_contents and p.style.name.startswith('Heading'):
            break
        if in_contents and p.text in targets:
            text = p.text
            p.clear()
            p.paragraph_format.tab_stops.add_tab_stop(Pt(485), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            link(p, text, targets[text]); p.add_run('\t'); field(p, 'PAGEREF '+targets[text]+' \\h')
    for p in headings:
        if p.text == '目錄': continue
        if p.style.name in {'Heading 1','Heading 2'}:
            nav = OxmlElement('w:p'); p._p.addnext(nav)
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(nav, p._parent)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(3)
            link(paragraph, '回目錄', contents_anchor)
            paragraph.add_run('  ｜  ')
            number = re.match(r'^(\d+(?:\.\d+)?)\s',p.text)
            number = number.group(1) if number else ''
            link(paragraph, '卡控條件', numbered.get(rules_for.get(number,''),rules_anchor))
            if number in cases_for:
                paragraph.add_run('  ｜  ')
                link(paragraph,'成功與失敗案例',numbered[cases_for[number]])
            if number in case_rules:
                paragraph.add_run('  ｜  ')
                link(paragraph,'錯誤與重試',numbered['17'])
            if number in {'12.3','5.6'}:
                for caption,key in [('報單案例','13.6'),('提單案例','13.7'),('Cost 案例','13.8')]:
                    paragraph.add_run('  ｜  ');link(paragraph,caption,numbered[key])
            if number == '15.2':
                paragraph.add_run('  ｜  ');link(paragraph,'NCPI 案例',numbered['13.3'])
            for r in paragraph.runs: r.font.size = Pt(8)
    # Explicit chapter tokens are shared by task tables, screenshots and error index.
    paragraphs=list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells: paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        text=paragraph.text
        if '[[' not in text: continue
        paragraph.clear();cursor=0
        for match in re.finditer(r'\[\[([\d.]+)\|([^\]]+)\]\]',text):
            paragraph.add_run(text[cursor:match.start()])
            if match.group(1) not in numbered: raise ValueError('Unknown manual target '+match.group(1))
            link(paragraph,match.group(2),numbered[match.group(1)])
            cursor=match.end()
        paragraph.add_run(text[cursor:])
    # Turn textual chapter references into working links as well.
    chapter_targets = {}
    for text, anchor in targets.items():
        m = re.match(r'^(\d+(?:\.\d+)?)\s', text)
        if m: chapter_targets[m.group(1)] = anchor
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if '詳細步驟' in p.text: continue
                    m = re.match(r'^請參閱第 ([\d.]+) 節',p.text)
                    if m and m.group(1) in chapter_targets:
                        text=p.text; p.clear(); link(p,text,chapter_targets[m.group(1)])
    for section in doc.sections:
        footer=section.footer.paragraphs[0]
        footer.clear(); footer.add_run('ICP 系統操作手冊 版本 1.5  ｜  第 ')
        field(footer,'PAGE'); footer.add_run(' 頁／共 '); field(footer,'NUMPAGES'); footer.add_run(' 頁')
    settings=doc.settings.element
    update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
