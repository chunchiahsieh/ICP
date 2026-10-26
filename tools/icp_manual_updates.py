"""Current user-facing workflows shared by the manual builder (2026-10-26)."""
import json
import re
import xml.etree.ElementTree as ET


def add_updates(key, doc, table, step, bullet, root):
    def heading(text):
        doc.add_heading(text, level=3)

    reminders = {
        '3.2': '完成判斷：驗證通過僅供預覽，須 Save 後查回 Ship Info。若 Save 中斷，先查是否已保存，再決定重試；不要任意改 Invoice 避開重複檢查。',
        '3.3': '更新前提醒：依 Invoice 更新所有符合的 Header；13 個更新欄位的空白可能清除原值。先核對完整資料再 Save；結果不明先查回資料，勿直接重送。',
        '3.4': '更新前提醒：此頁也使用共用 Invoice 更新流程；不要假設頁面名稱會自動限定所有目標。空白可能清除原值；Save 後重新查詢確認。',
        'settings': '整批取代提醒：先下載備份，核對完整清單；原有 A、B、C，本次只有 A、B 時，C 將不保留。上傳失敗應先修正再傳；結果不明先重查整頁。',
        '5.4': '覆蓋提醒：Save 依 Invoice 替換既有明細，非按 Material 合併。原有 5 筆、本次只提供 2 筆時，不能假設舊明細會保留。Save 後查正式批次與資料；網路中斷先查結果。',
        '5.6': '完成判斷：依上傳結果及清單查回資料；覆蓋確認前先核對 Invoice 與既有資料。確認過期或資料有變更時重新驗證。附件上傳與 Customs Data 是分開的操作。'
    }
    if key in reminders:
        heading('操作前與完成後核對')
        bullet(reminders[key])
        bullet('請參閱第 15.3 節失敗重試方式；操作結果不明時，先確認已保存內容再決定是否重試。')

    if key == 'common':
        heading('篩選方式與資料保存')
        table(['條件', '輸入與結果'], [
            ('Text', '一般文字欄位依頁面設定進行文字比對。'),
            ('多行 Text', '指定多選文字欄位每行輸入一個完整值。輸入 1、換行、2，結果為符合 1 或 2，並非包含 1 或 2 的文字。'),
            ('Checkbox', '可搜尋選項後勾選多個值；同欄選項為 OR，不同欄條件同時成立。需要空值時選取「空白」。'),
            ('DateRange', '設定開始及結束日期；Reset 清除條件。'),
            ('水平捲軸', '欄位超出畫面時，使用表格下方水平捲軸查看右側欄位。'),
            ('Upload 與 Save', 'Upload Check／預覽通過只代表可供檢查；仍須按 Save 才正式儲存。')
        ], [1.3, 5.7])
        bullet('Add DI/SA、Mass Update、Forwarder 需預覽後 Save。Settings Excel 與 Languages 匯入、Customs Data／附件上傳則依各頁確認直接處理；不要假設所有上傳都還有另一個 Save 步驟。Export 按上傳並通知 Hub 會建立產檔工作。')
    elif key == 'ship':
        heading('Deposit 與 ARUR 起案')
        for s in ['選取一筆 Header，核對 Invoice 與 Detail，再按 Deposit 或 ARUR。',
                  '在右側預覽核對表頭、明細與驗證訊息；Deposit 下方另有郵件預覽。',
                  '若畫面顯示缺值或長度錯誤，先取消預覽、修正資料，再重新開啟。',
                  '按提交後，在確認視窗再次核對資料並確認起案；關閉或取消不會建立案件。',
                  '返回清單檢查案件狀態。起案後的後續整合為背景處理，不以按鈕已點擊作為整合成功依據。']:
            step(s)
        bullet('ARUR 欄位名稱的黃色底色是重點標示，不代表這些欄位全部必填。缺值或驗證錯誤以畫面訊息為準。')
        bullet('ARUR 檢查 Receiver、Forklift、MovingLabor、WasteDisposal、DriverDetails；另檢查 Warehouse、TET PO 與起案主旨長度。')
        doc.add_page_break()
        heading('ARUR 表頭預覽排列')
        labels = {e.attrib['name'].replace('ShipInfo.Header.Field.', ''): e.findtext('value', '')
                  for e in ET.parse(root / 'ICP/Resources/SharedResource.zh-TW.resx').getroot().findall('data')
                  if e.attrib.get('name', '').startswith('ShipInfo.Header.Field.')}
        js = (root / 'ICP/wwwroot/js/shipinfo-case-drawer.js').read_text(encoding='utf-8-sig')
        names = re.findall(r"'([^']+)'", re.search(r'var arurHeaderOrder = \[([\s\S]*?)\];', js).group(1))
        highlight = set(re.findall(r"'([^']+)'", re.search(r'var arurHighlightedHeaderFields = new Set\(\[([\s\S]*?)\]\)', js).group(1)))
        rows = []
        for i in range(0, len(names), 3):
            rows.append(tuple(labels.get(n, n) for n in names[i:i+3]) + ('',) * (3-len(names[i:i+3])))
        t = table(['左側欄位', '中間欄位', '右側欄位'], rows, [2.33, 2.33, 2.34])
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        for i, n in enumerate(names):
            if n.lower() in highlight:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'FFFF00')
                t.rows[1+i//3].cells[i%3]._tc.get_or_add_tcPr().append(shd)
        doc.add_paragraph('上表為一般寬螢幕的三組欄位順序；窄螢幕依序排列成單組。顯示名稱依介面語言與設定而不同。')
        heading('通知信與寄送日期')
        table(['通知', '操作與核對重點'], [
            ('到貨通知', '在 Ship Info 的通知工具列開啟到貨通知排程；核對 Invoice、收件者、排程時間及預覽後，依畫面確認。查看排程狀態，避免重複建立。'),
            ('延遲交貨通知', '填寫延遲原因及 Delay Notification Date，開啟通知預覽。Delay Notification Date 為通知日期；郵件內 Delivery Date 使用 Delivery Information 的 Delivery Date。'),
            ('管制品到貨通知', '開啟管制品通知操作區，核對通知原因、收件者及預覽，再依畫面確認。'),
            ('收件者與權限', '可用按鈕受角色權限及資料狀態限制；排程或預覽通過不代表 SMTP 已寄送成功，應核對狀態與錯誤訊息。')
        ], [1.4, 5.6])
        bullet('到貨通知狀態依序可能為 Pending（待寄）、Sending（寄送中）、Sent（已寄）或 Failed（失敗）；只有 Pending 排程可使用取消按鈕。核對排程對應的郵件預覽後再決定是否取消。')
    elif key == '3.2':
        heading('新版範本與匯入步驟')
        for s in ['按 Download Sample 下載 DISA_Data Upload Template.xlsx，使用新版欄位，不自行增加系統管理欄位。',
                  '第一個分頁必須是 Upload Data，第一列為標題，第二列開始填資料。保留完整欄位名稱。',
                  '日期文字可填 yyyy-MM-dd 或 yyyy/MM/dd；例如 2026/09/12 會轉成 2026-09-12。編號含前導零時以文字保存。',
                  '按 Upload File，逐列查看驗證訊息；預覽使用水平捲軸核對完整欄位，不只查看畫面左側。',
                  '核對無誤後按 Save。若驗證失敗，依列號及欄位修正原檔後重新上傳。']:
            step(s)
        required = re.search(r'RequiredProperties\s*=\s*\[([\s\S]*?)\]', (root/'ICP/Helpers/AddDiSaImportRules.cs').read_text(encoding='utf-8-sig')).group(1)
        required_names = re.findall(r'"([^"]+)"', required)
        heading('必填與一致性卡控')
        doc.add_paragraph('必填 33 個欄位如下；欄位代碼以目前程式對應為準。')
        table(['類別', '必填欄位'], [
            ('表頭與運送', 'SA_DATE、INVOICE_NO、FORWARDER、BROKER、SHIPPER、ETD、ETA、INVOICE_DATE、MAWB、HAWB、FLT、FREIGHT、DESTINATION_PORT、DESTINATION_COUNTRY、WAREHOUSE、INVOICE_TYPE、INCOTERMS、ORDER_TYPE、DELIVERY_TO、BU、TOTAL_CARTONS'),
            ('明細', 'INVOICE_SEQ、ITEM_NO、DESCRIPTION、QTY、UOM、COO、PRICE、AMOUNT、CURRENCY、RATE、CARTON_NO、GROSS_WEIGHT')
        ], [1.3, 5.7])
        assert len(required_names) == 33
        bullet('ICP 已存在的 INVOICE_NO 不可再匯入；同一檔案的 INVOICE_NO + INVOICE_SEQ 不可重複。')
        bullet('同一 INVOICE_NO 只有以下明細欄位允許不同：INVOICE_SEQ、ITEM_NO、DESCRIPTION、QTY、UOM、COO、PRICE、AMOUNT、DELIVERY_LINE_NO、ECCN、EL_FLAG、SDS_FLAG、HAZMAT、WBS_ELEMENT、NET_WEIGHT_OF_THE_ITEM。其餘上傳欄位須一致。')
        bullet('數值欄位必須能解析為數字；選填數值欄有填值時仍會檢查。MDP_FLAG 只可填 Y 或空白。WBS_ELEMENT 保存至明細。')
        table(['上傳欄位', 'Settings 比對'], [('BROKER','Broker Key'),('SHIPPER','Shipper Key'),('FREIGHT','Air/Sea Key'),('WAREHOUSE','Warehouse Codes Key'),('INVOICE_TYPE','Invoice Type Key'),('ORDER_TYPE','Order Type Key'),('DELIVERY_TO','Delivery To List 的 Customer Code，非 Key 或 Name'),('BU','BU Codes Key')], [2.0, 5.0])
    elif key == '3.5':
        heading('Case Mark 輸出核對')
        bullet('產生 Case Mark 後，開啟檔案核對 CASE MARK、Made in、字體、行距及箱號。列印時檢查 A4 預覽及縮放，避免印表機設定造成字體過小或裁切。')
        bullet('等待產生中時先確認狀態，不要連續建立同一份匯出；若失敗，保留錯誤訊息交管理者檢查。')
    elif key == 'settings':
        heading('各頁 Excel 下載與整批取代')
        for s in ['進入要維護的 Settings 頁面，按 Download Excel 下載該頁目前有效資料。下載需具有查看該頁的權限。',
                  '保留下載檔的欄位名稱與結構，在該頁專用資料中修改；不要把其他 Settings 頁的資料混入。',
                  '按 Upload Excel 選取 .xlsx 檔，需具有該頁新增資料權限。',
                  '確認整批取代訊息。系統先驗證所有列；全部通過才以本檔資料取代該頁原有資料，不是合併新增。',
                  '失敗時原資料保留，依列號與錯誤修正後重傳；成功後重新查詢，必要時再下載核對。']:
            step(s)
        bullet('下載不受目前畫面篩選限制，取得該頁範圍的有效設定。空白資料檔不可用來清空整頁。公式、重複設定或不符欄位格式會遭拒絕。')
        bullet('Shipper、Country of Origin 與其他 Settings 頁分別維護；各頁上傳只取代各自的範圍。Delivery To List 的 Customer Code 與 Name 是不同欄位，請勿互換。')
    elif key == '5.1':
        heading('Header 與 Detail 篩選')
        bullet('Shipping Report 保留 Header 與 Detail。先在 Header 篩選及選取資料，再核對對應明細；下載前確認目前條件。')
        bullet('Invoice No、MAWB、HAWB、TET PO、NCDR No、End User 支援多行精確篩選，每行一個值。一般文字欄位仍依其設定比對。')
        bullet('Checkbox 包含「空白」選項；日期欄位使用 DateRange。數字靠右，顯示小數位依欄位設定，原資料精度不因畫面顯示而改變。')
        report_fields(root, 'shipping-report-table-fields.json', table, doc)
    elif key == '5.3':
        heading('僅保留 Header')
        bullet('本頁已移除 Detail，也不需選取 Header 載入明細。Download Excel 與 Header 顯示欄位一致，一筆 Header 對應一筆輸出。')
        bullet('Invoice No、NCDR No、MAWB、HAWB 可輸入多行完整值。SaDate、ETA、DeliveryDate、DelayNotificationDate 使用日期區間；預設依建立時間由新到舊排列。')
        report_fields(root, 'mass-data-report-table-fields.json', table, doc)
    elif key == '5.4':
        heading('Upload Check')
        for s in ['按 Download Sample 取得 ForwarderDataUploadTemplate.xlsx。第一分頁名稱須為 Upload Data，第一列為欄位名稱；資料由第二列開始。填寫範例分頁不會匯入。',
                  '在 Upload Check 上傳檔案，核對預覽與錯誤；日期文字接受 yyyy/MM/dd 或 yyyy-MM-dd，顯示為 yyyy-MM-dd。',
                  '通過後按 Save；若有既有資料覆蓋確認，核對影響範圍後才繼續。未儲存的預覽不會列為已匯入批次。']:
            step(s)
        table(['規則', '說明'], [('選填欄位','只有 Customer Reference、FLIGHT 2#、Action、MDP 四欄選填，其餘 18 欄必填。'),('數量與日期','Quantity 必須是數字；Confirmed Custom Date、ETD、ETA 為日期。'),('C/B','必須存在於 Settings → Broker 的 Key，不使用顯示名稱代替。')], [1.4,5.6])
        heading('Upload History 與原檔下載')
        for s in ['切換 Upload History，按 Refresh 更新已儲存成功的批次。',
                  '可依檔名、匯入時間、匯入人、筆數、狀態及原檔可用情況進行欄位篩選；實際可見欄位由設定決定。',
                  '按 Download original 下載該批原始檔。已被後來資料覆蓋的舊批次仍保留紀錄與可用原檔。',
                  '若顯示 File not found，代表紀錄存在但原檔目前不可取得，請聯絡管理者檢查保存路徑；不是要求重新儲存同一批。']:
            step(s)
        bullet('檔案選取、驗證失敗與未按 Save 的預覽，不算成功匯入紀錄。')
    elif key == '5.6':
        heading('Customs Data 上傳與欄位規則')
        for s in ['按 Download Template 下載 CustomsDataUploadTemplate.xlsx，第一個分頁為 Upload Data，第一列標題，第二列起資料。',
                  '按 Upload Customs Data，依錯誤訊息檢查必填、日期、數字及設定值。若系統要求覆蓋既有資料，先核對再確認。',
                  '成功後在清單核對 HAWB、報單、日期與金額，使用欄位篩選或 Download Excel。']:
            step(s)
        table(['欄位', '目前規則'], [('CIF Value','選填。'),('Declaration Amount (TWD)','必填，且必須為數字。'),('Air/Sea','必填，需符合 Settings 的 Air/Sea 設定。'),('Country of Origin','篩選選項對應 Settings → Country of Origin。'),('Exchange Rate','固定顯示 4 位小數。'),('Net Weight (KG)','固定顯示 3 位小數。'),('Unit Value','固定顯示 2 位小數。')], [2.0,5.0])
        bullet('固定顯示位數不等於限制匯入的小數位數；不得以畫面已四捨五入，推定原始數值已被截斷。頁面路由仍為 /TariffData。')
        heading('Declaration PDF、Import Doc 與 Cost')
        table(['按鈕／類別','檔案規則'], [('Upload Declaration PDF','報單：檔名為 HAWB；支援 .pdf、.tif。'),('Upload Import Doc','提單：檔名為 HAWB；支援 .pdf、.xls、.xlsx。'),('Upload Cost','費用檔維持 Excel 上傳流程，與提單／報單分開操作。')], [2.2,4.8])
        for s in ['確認要關聯的 HAWB 已存在，依類型選擇對應上傳按鈕。',
                  '上傳完成後，從 Declaration PDF 或 Import Doc 欄位開啟附件清單，核對檔名及版本並下載。',
                  '同一 HAWB 的資料會共用對應附件。替換檔案時保留版本；刪除目前版本前先核對，系統依可用版本更新關聯。']:
            step(s)


def report_fields(root, filename, table, doc):
    config = json.loads((root/'ICP/Config'/filename).read_text(encoding='utf-8-sig'))
    for area in ('header', 'detail'):
        fields = config.get(area, {}).get('list', {}).get('fields', [])
        fields = [f for f in fields if f.get('visible', True)]
        if not fields:
            continue
        doc.add_paragraph(('Header' if area == 'header' else 'Detail') + ' 欄位順序與篩選')
        rows = [(str(i), f['fieldName'], ('多行精確比對' if f.get('multiValueText') else f.get('filterType','Text')) if f.get('searchable',False) else '不提供篩選') for i,f in enumerate(fields,1)]
        table(['順序','欄位代碼','篩選'], rows, [0.6,3.5,2.9])
