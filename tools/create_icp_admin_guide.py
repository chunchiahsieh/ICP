"""Build the administrator and developer guide from current source inventories."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from icp_manual_navigation import bookmark, link, field
from icp_admin_playbook import EXTRA_CHAPTERS, add_reading_key, append_playbook
from icp_admin_findability import quick_index, change_matrix, maintenance_record
from icp_admin_completion import keywords, log_examples, doc_update_checklist, make_templates

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/manual'
DOCX=OUT/'ICP系統管理者手冊-v2.4.docx'
doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Microsoft JhengHei';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
for n,size in [('Title',27),('Heading 1',18),('Heading 2',13)]:
    doc.styles[n].font.size=Pt(size);doc.styles[n].font.bold=True
title_props=doc.styles['Title']._element.get_or_add_pPr()
for border in list(title_props.findall(qn('w:pBdr'))):title_props.remove(border)

in_body=False
current_chapter=0
def p(t,kind=None):
    if t and in_body:
        if kind is None:
            kind='維運建議' if t.startswith(('接手時','回復先','上列是','測試程式位於','實際 Git','IIS 維運建議','官方參考')) else '目前程式行為'
        t='【'+kind+'】'+t
    para=doc.add_paragraph(t)
    if not t:
        para.paragraph_format.space_after=Pt(2)
        para.paragraph_format.line_spacing=Pt(2)
    return para
def b(t): return doc.add_paragraph('【維運建議】'+t,style='List Bullet')
def h(t): return doc.add_heading(t,level=2)
def table(head, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(head));t.autofit=False
    if widths is None: widths=[7/len(head)]*len(head)
    for i,x in enumerate(head):
        c=t.rows[0].cells[i];c.text=x
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F4E78');c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            cells[i].text=str(x)
            if j%2:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F3F6F9');cells[i]._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                if len(rows)<=3:para.paragraph_format.keep_with_next=True
                for run in para.runs:run.font.size=Pt(9.5)
    for c in t.rows[0].cells:
        for para in c.paragraphs:para.paragraph_format.keep_with_next=True
    p('');return t
def code(lines):
    if lines and lines[0].startswith(('[INF]','[WRN]','[ERR]')):
        doc.add_paragraph('【維運建議】教學用 Log 格式示例；不可當作指令執行或本次實測紀錄。')
    else:
        doc.add_paragraph('【目前程式行為】流程或 API 摘要。' if current_chapter in {2,10,11} else '【維運建議】以下是待執行範例，先核對環境與占位符。')
    for line in lines:
        para=doc.add_paragraph(line);para.paragraph_format.space_after=Pt(2)
        for r in para.runs:r.font.name='Consolas';r.font.size=Pt(8.5)

chapters=[
 '1 手冊定位與快速接手','2 系統架構與資料流','3 開發環境與建置','4 程式分層與修改入口',
 '5 設定檔與環境選用','6 畫面欄位與共用元件','7 資料庫與版本異動','8 匯入與附件開發',
 '9 報表與資料權限','10 起案與整合事件','11 Export 與產檔服務','12 通知信開發與管理',
 '13 IIS 與背景服務維運','14 部署與回復流程','15 問題排查手冊','16 測試與驗收',
 '17 擴充功能與維護規範','18 程式索引與版本紀錄']+EXTRA_CHAPTERS
doc.add_paragraph('ICP系統管理者手冊',style='Title')
doc.add_paragraph('開發者與維運人員指南',style='Subtitle')
p('版本 2.4　更新日期 2026-10-27')
p('本手冊由原「ICP 系統設定管理手冊」擴充，協助開發者理解程式、修改功能、建置部署及追查問題。一般使用者操作另見 ICP 系統操作手冊 v1.5。')
p('涵蓋 ICP、TEL.IntegrationHub 與 ICPFileGenerator，以本地端測試作為完成依據，不設正式環境驗收表。文件不列真實密碼、連線字串與收件名單。')
doc.add_heading('目錄',level=1)
for title in chapters:
    entry=p(title);entry.paragraph_format.space_after=Pt(0);entry.paragraph_format.line_spacing=1.0
    for run in entry.runs:run.font.size=Pt(9.5)

def chapter(n):
    global in_body,current_chapter
    in_body=True
    current_chapter=n
    doc.add_heading(chapters[n-1],level=1).paragraph_format.page_break_before=True

chapter(1)
p('接手時先掌握執行環境、三個服務的設定來源、資料庫與共用檔案路徑，再依問題定位到 Controller、Service、Repository 或 Worker。先讀取狀態及紀錄，避免用重送業務資料診斷問題。')
p('常見問題與管理者／開發者閱讀入口集中於 [[1.3]]；先定位問題再跳到完整流程，避免只靠服務已啟動或 API 可回應判斷業務完成。',kind='維運建議')
p('文件中的程式路徑以儲存庫根目錄為基準。目前根目錄：'+str(ROOT)+'。部署路徑由主機實際站台設定決定，不應照抄開發機路徑。')
h('1.1 管理邊界')
b('本版只更新文件、下載來源與文件連結；未修改業務邏輯、未執行 SQL、未正式寄信或重送事件。')
add_reading_key(doc,table)
quick_index(table,p,h)
keywords(table,p,h)

chapter(2)
table(['服務','目前程式責任','依賴'],[
 ('ICP / net9.0','MVC、匯入、報表、權限、通知信、起案 Outbox、Export Request','SQL Server、身分與組織資料、RabbitMQ、SMTP、Hub、檔案目錄'),
 ('TELIntegrationHub / net9.0','起案 Consumer、MESSAGE_LOG、ILC 寫入、Outbox 完成回寫、Export HTTP 協調','Hub DB、ICP DB、ILC／相關資料來源、RabbitMQ'),
 ('ICPFileGenerator / net8.0','掃描與 Claim Job、產 Pickup Notice Excel／Case Mark PDF、通知 Hub 結果','與 ICP 同環境的業務資料庫、共用輸入／輸出、Hub HTTP')
],[1.65,3.3,2.05])
h('2.1 兩條整合路徑')
code(['起案：ICP → INTEGRATION_EVENT_OUTBOX → RabbitMQ → Hub Consumer',
      '      → ILC／目的資料 → Hub MESSAGE_LOG 與 ICP 完成回寫',
      'Export：ICP 上傳 → EXPORT_REQUEST → HTTP Hub → ICPFileGeneratorJob',
      '        → FileGenerationWorker → Excel／PDF → HTTP Hub → EXPORT_REQUEST'])
p('不要把 Export HTTP 工作流程與起案 RabbitMQ Outbox 混為一談。現有 README 的部分 Phase 1／Stub 描述已過時；本版以 Program.cs、Controller、Service 與 Repository 的實際路徑為準。')
h('2.2 主要儲存與識別')
table(['識別／儲存','用途'],[
 ('Invoice No、TET PO、Invoice Seq','業務定位；不同匯入流程的分組與唯一性規則不同，不能共用一個推測的資料鍵。'),
 ('messageId／Outbox Id','追查起案發布、Consumer 處理、MESSAGE_LOG 與回寫。'),
 ('RequestId／JobId','追查 Export Request、產檔 Job、檔案路徑與 Hub 回報。'),
 ('資料庫與檔案','資料保存與原檔保存是兩種狀態；原檔遺失不代表資料匯入失敗。')
],[2.0,5.0])

chapter(3)
p('以三個 .csproj 的 TargetFramework 與 PackageReference 為建置依據。ICP／Hub 需要 net9.0 對應環境，FileGenerator 為 net8.0；部署時也必須具備相應執行環境，不能只安裝其中一個就假設全部可用。')
table(['專案','重要套件與作用'],[
 ('ICP','EF Core SQL Server、Negotiate、ClosedXML、ExcelDataReader、RabbitMQ.Client、Serilog。'),
 ('TELIntegrationHub','MassTransit.RabbitMQ、EF Core SQL Server、Serilog、Swagger。'),
 ('ICPFileGenerator','Microsoft.Data.SqlClient、ClosedXML、QuestPDF、Swagger。')
],[1.7,5.3])
code(['dotnet --info','dotnet restore ICP.sln','dotnet build ICP.sln -c Release',
      'dotnet run --project ICP/ICP.csproj',
      'dotnet run --project TELIntegrationHub/TEL.IntegrationHub.csproj',
      'dotnet run --project ICPFileGenerator/ICPFileGenerator.csproj'])
p('啟動並非純讀取：ICP／Hub 啟動會執行部分資料表確保邏輯；ICP 同時啟動通知與 Outbox Worker。開發前先確認環境選用、資料庫、寄信啟用與整合目的地。不要把正式設定直接拿來啟動本機除錯。')
h('3.1 Visual Studio 與執行方式')
b('開啟 ICP.sln，依測試需要設定啟始專案。VS 停止除錯會結束相關處理序；結束碼 -1 本身不足以證明應用程式當機。')
b('IIS Express／launchSettings 的 localhost Port 是開發啟動設定，不是主機 IIS 的站台 Binding。主機上用 VS 開啟仍需檢查實際 Process、環境與設定來源。')
b('先單獨確認各服務啟動與資料連線，再進行跨服務業務測試；正式匯入、起案或寄信應使用核准的測試資料與環境。')

chapter(4)
table(['層次','位置','修改原則'],[
 ('HTTP／頁面入口','ICP/Controllers；Controllers/Setting','模型綁定、權限與反偽造、回應；不要只在 JavaScript 卡控。'),
 ('業務與匯入','ICP/Services、Helpers','驗證、欄位對應、交易、排程及整合；避免把同一規則散在多頁。'),
 ('資料存取','ICP/Repositories、Data/ApplicationDbContext.cs','欄位對映、查詢、DataScope、寫入與交易。'),
 ('畫面與共用元件','ICP/Views、wwwroot/js、wwwroot/css','欄位顯示、事件、過濾、捲軸、錯誤呈現與快取。'),
 ('資料契約','ICP/Models、Hub/Models、FileGenerator/Models','前後端 JSON、事件、DB 欄位與匯出資料需同步。'),
 ('可設定欄位與範本','ICP/Config、Files、Resources','設定與程式一同檢查發布；example 只作參考，不等於目前執行設定。')
],[1.4,2.8,2.8])
h('4.1 從需求找程式')
table(['需求','主要入口'],[
 ('DI/SA 規則','AddDiSaController → AddDiSaImportService → AddDiSaExcelColumnMap／AddDiSaImportRules'),
 ('Forwarder 與歷史','ForwarderDataUploadController、ForwarderDataImportService、ForwarderImportRules、ForwarderHistoryTableMetadataProvider'),
 ('Customs Data 與附件','TariffDataController、TariffDataImportService、TariffCustomsImportRules、TariffExcelColumnMap'),
 ('報表','ReportControllerBase、ReportMetadataProvider、report-data.js、Config/*report-table-fields.json'),
 ('Ship Info 與起案','ShipInfoController、ShipInfoRepository、ShipInfoCaseSnapshotBuilder、shipinfo-case-drawer.js'),
 ('Settings Excel','Controllers/Setting/SystemConfigControllerBase.cs、Helpers/SettingsExcelHelper.cs')
],[1.55,5.45])

chapter(5)
h('5.1 實際載入規則')
p('三個 Program.cs 都會先判斷機器名。AGA-PC 保留 WebApplication.CreateBuilder 的預設組態來源；非 AGA-PC 會清空來源，載入選定 TEL JSON，再載入環境變數及命令列。因此非 AGA 主機不會自動疊加 appsettings.json／Development。')
table(['機器條件','JSON 來源','核對方法'],[
 ('AGA-PC','預設來源；Development 時另含 appsettings.Development.json','看啟動 Log 的 AppSettings；並核對環境變數／命令列覆寫。'),
 ('TETIS87181','appsettings.TEL.TETIS87181.json','檔案為必要檔；不能只修改 appsettings.TEL.json。'),
 ('其他非 AGA-PC','appsettings.TEL.json','例如另一台開發機也會選 TEL 設定，不以 VS 是否 Debug 決定。')
],[1.6,2.7,2.7])
p('環境變數可用雙底線表示階層，例如 NotificationMail__ArrivalNotice__Enabled。陣列覆寫應核對每個索引，不要假設寫一個空陣列就一定清除所有來源的項目。')
h('5.2 設定字典')
table(['區段','主要設定與維護重點'],[
 ('App','BrandName、TitleSuffix、Mode、SuperUser、SimulatedWindowsIdentity、DevUser：品牌與登入／模擬身分。正式環境不得誤用開發模擬身分。'),
    ('ConnectionStrings','ICP：ICP_Connection、ILC_Connection、FIESTA_Connection。FileGenerator：ICPFileGenerator。Hub：HubDatabase、ICP_Connection、ICPFileGenerator、ILC_Connection、FIESTA_Connection；不同鍵需指向正確用途與環境。'),
 ('Sidebar','RecentEnabled、RecentLimit、PinnedLimit、DefaultPage、MenuOrder、CustomLinks。CustomLinks 的 Name／Url／Enabled／OpenInNewTab／SortOrder 控制文件入口。'),
 ('ForwarderDataUpload／TariffData','StoragePath、MaxSizeMb；TariffDataOptions 目前未定義舊文件所列 BrokerKeywords，勿把殘留設定當有效功能。'),
 ('ArurAttachment','IpcStorageRoot、RelativeFolder、MaxSizeMb、AllowedExtensions：核對實際附件服務讀取處及可用路徑。'),
 ('Integration','RabbitMq：Enabled、HostName、Port、VirtualHost、UserName、Password、Exchange、RoutingKey。Outbox：PollIntervalSeconds、MaxRetryCount、BatchSize。Hub：BaseUrl。Export：InputDirectory／OutputDirectory。'),
 ('NotificationMail','Smtp、DeliveryDelay、ArrivalNotice、ControlledGoodsArrival、DepositCase、ShipInfoUrl；明細見第 12 章。'),
 ('FileGenerator','PollingIntervalSeconds、WorkerId、OutputDirectory、ProcessingTimeoutMinutes、MaxRetryCount、Hub.BaseUrl；目前 Worker 使用 IOptions，修改後須重啟生效。'),
 ('LocalizationManagement','Enabled 控制語系管理功能；Resources 的部署及寫入權限需同步確認。'),
 ('Logging／Serilog／AllowedHosts','記錄層級、輸出與 Host 限制；實際檔案 sink 與 console 設定依各 Program.cs，不保證三服務完全相同。')
],[1.7,5.3])
h('5.3 變更與秘密管理')
b('實際環境設定被 Git 忽略，須由環境設定流程另行更新；不要為部署方便強制提交密碼或完整連線字串。新增鍵同步維護各相關環境及 example。')
b('改設定前保存受保護備份；改後核對載入來源、型別與完整值。reloadOnChange 不等於所有 IOptions 消費者都會熱更新；正式變更以重新啟動及驗證結果為準。')
b('文件範例使用占位符；不要把正式設定原文附到公開問題單或本手冊。')

chapter(6)
p('清單欄位、表單欄位及後端資料來源是三層契約。Config 能控制既有欄位的順序、可見與過濾，但新增 DB 欄位不能只加一筆 JSON。')
table(['檔案群','修改時需同步'],[
 ('shipinfo-pro-datatable-fields.json','Header／Detail 清單、ShipInfoFieldCatalog、DTO、Mapper、Repository。'),
 ('shipinfo-header-form-fields.json／shipinfo-detail-form-fields.json','表單 Metadata、可編輯欄位、查詢選項、後端驗證與儲存。'),
 ('shipping-report／mass-data-report／compare-icp-vs-arur-table-fields.json','各報表 Metadata、資料服務、畫面及 Excel 欄位。'),
 ('forwarder-table-fields.json／forwarder-upload-history-table-fields.json','現行資料與歷史批次是分開設定；歷史日期、狀態、原檔下載另查對應 Provider／JS。'),
 ('tariff-table-fields.json／customs-data-download-table-fields.json','顯示、數值對齊、精度、篩選選項來源及匯出。')
],[3.2,3.8])
h('6.1 Filter 與表格維護')
b('多行 Text 每行一個完整值，以 OR／IN 語意比對；不同欄條件同時成立。一般 Text 的行為依欄位設定，不可全部一律改成 IN。')
b('Checkbox 必須包含目前資料的全部可用樣態及空白；來源為 Settings 時，須核對 Category 及 Key／Value，不能只更改標籤。DateRange 要同時確認查詢邊界、日期儲存格式及時區。')
b('數值靠右與固定顯示小數是呈現規則，不代表上傳時截斷原值。報表畫面及 Excel 的資料範圍應一致。')
b('共用 pro-datatables.js、pro-table-filters.js 的修改要回歸不同報表及資料列很少的情況。檢查最小高度、水平捲軸、sticky、overflow 與 dropdown 層級，不能只調 z-index 就認定裁切已修復。')
h('6.2 新增欄位實作順序')
for t in ['確認欄位歸屬 Header／Detail、資料型別、空值與最大長度。','先備妥 SQL／Entity mapping，再更新 DTO、Mapper、Catalog／Metadata。','同步匯入映射與驗證、編輯保存、清單／篩選、匯出及語系。','增加邊界案例並檢查實際 publish 輸出含 Config、Resources、Files。']:b(t)

chapter(7)
p('ICP 的 EF 對映集中於 Data/ApplicationDbContext.cs；Hub 與 FileGenerator 亦直接使用 ICP 業務表。資料庫名稱依環境連線字串，不能一律假設都叫 TEL-ICP 或 ICP。')
table(['資料群','主要表','追查用途'],[
 ('業務','ICP_HEADER、ICP_DETAIL、TariffData、ForwarderDataUpload、ForwarderDataUploadArchive','Invoice、明細、報關資料與歷史批次。'),
 ('設定／權限','SystemConfigs、Resources、Roles、RolePermissions、RolesTELID、RolesDepID、RolesMailGroup','Settings 分頁、資源與角色授權；軟刪除及唯一鍵需一併核對。'),
 ('稽核／附件','Attachments、SHIPINFO_AUDIT_LOG、DELETE_LOG','檔案實體與紀錄、編輯及刪除線索。'),
 ('整合／產檔','INTEGRATION_EVENT_OUTBOX、EXPORT_REQUEST、ICPFileGeneratorJob、Hub MESSAGE_LOG','事件、重試、請求與產檔狀態。'),
 ('通知','NOTIFICATION_MAIL_DISPATCH、NOTIFICATION_MAIL_LOG、ARRIVAL_NOTICE_SCHEDULE、CONTROLLED_GOODS_NOTICE_SCHEDULE','排程、寄送與失敗內容。')
],[1.1,3.6,2.3])
h('7.1 Schema 發布原則')
p('ICP 啟動會確保部分稽核、Outbox、Export Request 與 Forwarder Archive 結構，但不代表每個新增欄位都會自動補齊。不要用程式已啟動推論所有部署 SQL 已執行。')
table(['腳本','適用位置與順序'],[
 ('ICP/sql/017_AddDetailWbsElement.sql','先在 ICP 業務庫執行再部署相依程式；新增 Detail WBS_ELEMENT，並以 Invoice＋TET PO 回填既有 Header 值。動態 SQL 避免同批次編譯看不到新增欄位。'),
 ('ICP/sql/018_AddTariffImportDoc.sql','在 ICP 業務庫新增 TariffData.ImportDoc nvarchar(500)，既有列為 NULL；檔案不能是空檔。'),
 ('ICPFileGenerator/sql/001、002','核對 Job 表及 Claim procedure；同環境 ICP 資料庫。003 為示例，不應當成正式部署步驟執行。'),
 ('TELIntegrationHub/sql/004_DropRoutingRule.sql','先部署移除依賴的 Hub，再備份並按核准流程刪除舊表；腳本含 USE TEL_IntegrationHub，執行前核對目標。現有 HubDbContext 已無 ROUTING_RULE entity。')
],[2.65,4.35])
b('只執行本版必要 SQL 並記錄環境、雜湊及結果；破壞性變更先確認舊版相依。')
b('程式回復不會回復資料；依備份與核准腳本處理 schema／資料，詳見第 14、22 章。')

chapter(8)
table(['功能','規則與實作契約'],[
 ('Add DI/SA','第一頁 Upload Data、第二列起；73 欄的映射與卡控見 Helpers。日期接受兩種分隔格式；Invoice 已存在及 Invoice＋Seq 重複會阻擋，同 Invoice 只有指定 Detail 欄位可不同。'),
 ('Mass Update','兩頁共用 MassUpdateImportService；以 Invoice 定位，空白可能清除 13 個更新欄位。不能假設畫面分類就是更新範圍限制。'),
 ('Forwarder','22 欄中 18 必填，僅 Customer Reference、FLIGHT 2#、Action、MDP 選填。C/B 比 Broker Key；同 Invoice 的 MAWB／HAWB／FLIGHT 1#／FLIGHT 2#／ETA 要一致。先驗證，Save 時替換 Invoice 資料並保留成功批次原檔。'),
 ('Customs Data','CIF 選填；Declaration Amount(TWD) 必填且解析數字；Air/Sea、Shipper 比對設定 Value。Broker 必填不等於也有 Broker Key 存在檢查。COO 選項來源與上傳驗證不能混淆。'),
 ('Settings Excel','各頁獨立範圍；Download 同 View，Upload 同 Create。先驗證全部列，再以 Serializable 交易將原作用範圍軟刪除並加入新資料；不是跨所有 Settings 頁全刪。'),
 ('Customs 附件','Declaration PDF：pdf／tif；Import Doc：pdf／xls／xlsx；以 HAWB 命名並對應既有資料。Cost 的 Excel 是另一路徑。')
],[1.4,5.6])
h('8.1 修改卡控的核對清單')
b('驗證服務須保持可測試，不依賴前端必填星號。錯誤需含列號、欄位與可修正原因；回傳換行字元時，前端需用可顯示換行的方式安全呈現。')
b('保留數字編號前導零；日期格式、decimal 精度、必要／可選及最大長度依各功能規則，不用一個寬鬆 Parse 套全部欄位。')
b('資料庫交易與原檔保存皆需測試失敗情境；資料成功而原檔不可下載、檔案成功而 DB 失敗，應有明確結果與補救方式。')
b('範本來源為 ICP/Files，新增／換檔要核對 csproj 的發布規則。測試應使用實際最新版表頭，不能只對 helper 的硬編碼欄位做鏡像測試。')

chapter(9)
h('9.1 查詢與匯出')
p('Shipping Report 保留 Header／Detail；Mass Data Report 只保留 Header，下載欄位應同 Header。ReportControllerBase 提供共用查詢／匯出入口，Metadata 決定欄位，資料服務與 Repository 負責過濾。')
b('新增 Filter 時同步 Request model、欄位選項、前端保存條件、後端查詢及下載，不可只換 Config 類型而留下原 Text 查詢。')
b('Checkbox 的全域選項查詢不能只查目前一頁資料；空白值與查詢條件需一致。文字多值應參數化，勿串接使用者輸入形成 SQL IN。')
h('9.2 身分與授權')
table(['元件','功能與維護'],[
 ('WindowsIdentityMiddleware／IisWindowsIdentityHelper','辨識實際使用者，結合 App.Mode／模擬身分；先核對登入帳號再查業務權限。'),
 ('UserResourcePermissionService／PermissionScannerService','資源權限與掃描、角色指派、Session 權限。新增 action 不只要改側欄，也須有後端授權。'),
 ('PageDataScopeService','依頁面 View resource，在分頁及匯出前套用範圍；Detail 透過 Header Invoice＋TET PO 關聯。另有 CustomsDataDownload 專用範圍服務。'),
 ('RolePermissions／Session','修改權限後重新載入權限或登入；不要把舊 Session 的結果當成新設定無效。')
],[2.9,4.1])
p('重要現況：PageDataScopeService 的註解與實作明確採用「缺少或無效的 scope 視為不限制」策略。這不是 fail closed；管理者應驗證 JSON、條件與使用者測試結果，不能以填了某段 JSON 就推論資料已限制。多授權範圍亦應驗證合併結果。')

chapter(10)
p('ShipInfoCaseSnapshotBuilder／EventFactory 組合起案快照與事件，Outbox Publisher 依 Integration.Outbox 輪詢；RabbitMq.Enabled=false 時事件保留待發。')
table(['階段','檢查內容'],[
 ('建立快照','Deposit 與 ARUR 的必填不同；核對 Header／Detail、主旨、附件、收件預覽與業務欄位。黃底只表示重點欄位，不代表全部必填。'),
 ('Outbox 發布','Id／messageId、PayloadJson、重試次數、錯誤與狀態；Published 不等於下游已完成。'),
 ('Hub Consumer','DepositCaseInitiatedConsumer／ArurCaseInitiatedConsumer；檢查 MESSAGE_LOG 與相依 ILC 資料寫入。'),
 ('完成回寫','IcpOutboxCompletionService 以識別資訊回寫 ICP；核對回寫目標庫與目前案件狀態。')
],[1.55,5.45])
h('10.1 API 與重試限制')
code(['GET <HUB_BASE>/health','GET <HUB_BASE>/api/messages',
      'GET <HUB_BASE>/api/messages/errors','GET <HUB_BASE>/api/messages/{messageId}'])
p('MessagesController 的 POST /api/messages/{messageId}/retry 目前回傳 501 未實作，不能寫入操作手冊當作可用恢復按鈕。人工重送前，先查下游是否已寫入及重複保護，並保留事件內容、目的系統與處理紀錄。',kind='尚未實作')
p('ROUTING_RULE 已自現行 Hub EF 對映移除；目前依 Consumer／程式路徑處理，新增目標需修改契約、註冊與處理服務，不應再教管理者新增舊 routing table 資料。')

chapter(11)
h('11.1 HTTP 協調與檔案')
table(['入口','作用'],[
 ('Hub POST /api/export/export-requests','接受 Request，將 EXPORT_REQUEST 標為 Processing 並建立 Job；Request 必須存在於相同環境 ICP DB。'),
 ('FileGenerationWorker','先處理逾時，再 Claim Job，呼叫 FileGenerationService；成功保存 Job 結果後回報 Hub。'),
 ('Hub POST /api/export/file-jobs/completed 或 failed','回報產檔結果並更新 EXPORT_REQUEST。這些是變更狀態的 API，不能隨意用正式 Request 測試。'),
 ('ShippingAdviceSheetReader／產檔器','讀取 Shipping advice；PickupNoticeExcelGenerator 與 CaseMarkPdfGenerator 產生輸出，查 CountryOfOrigin／PickUpLocation。')
],[3.2,3.8])
p('ICP Integration.Export.InputDirectory／OutputDirectory、FileGenerator.OutputDirectory 的預設相對路徑不保證在不同 ContentRoot 下相同。應換算成實體完整路徑，確認輸入可讀、輸出可寫，以及 ICP 下載可讀的是同一位置。跨機時須共用路徑與正確服務帳號權限。')
h('11.2 唯讀診斷 API')
code(['GET <FILEGEN_BASE>/api/system/status',
      'GET <FILEGEN_BASE>/api/file-generation-jobs?status=Processing',
      'GET <FILEGEN_BASE>/api/file-generation-jobs/{id}',
      'GET <FILEGEN_BASE>/api/file-generation-jobs/request/{requestId}'])
p('status 回 Running 只是 API 處理序可回應；目前 SystemController 回傳固定 version=1.0.0，沒有查 DB、Worker 心跳或 Hub 連線。不能用此欄位確認實際發布版本，也不能以 status 200 判定產檔鏈健康。')
h('11.3 失敗恢復')
b('Job Completed 而 ICP 仍 Processing：查 Hub 回報、RequestId 與連線庫，勿先刪 Job 或重建 Request。')
b('逾時處理依 ProcessingTimeoutMinutes／MaxRetryCount、JobRepository 與 Claim SQL；不保證所有 Failed 自動重試。')

chapter(12)
table(['程式／資料','責任'],[
 ('NotificationMailOptions','Smtp 與四類通知；各類有 Enabled、Frequency、SendTimes、Title、MailTo、CcTo。'),
 ('NotificationSmtpClient','共用 SmtpClient；Host=smtp.tel.co.jp、Port=25、EnableSsl=false 為目前預設／核准設定。FromAddress 非空保留，缺少時使用 TETIS87GEMSAdmin@tel.com。'),
 ('NotificationMailSender','先寫 NOTIFICATION_MAIL_LOG 為 Sending，再送信；失敗寫 Failed／ErrorMessage，成功寫 Sent／SentUtc。寄送結果與收件信箱投遞是不同階段。'),
 ('Workers／ScheduleService','DeliveryDelay、ArrivalNotice、ControlledGoodsArrival 背景處理；DepositCaseMailService 為押金通知流程。排程服務管理 Pending 等狀態。')
],[2.3,4.7])
h('12.1 收件者與時段變更程序')
p('例：新增某類通知的副本收件者。管理者先記錄環境、通知種類、原名單、新名單、生效時間及原因；確認實際 AppSettings 來源後，只修改該種類 CcTo 陣列，沿用既有格式。相關生效影響見 [[20.4]]。',kind='維運建議')
for t in ['ArrivalNotice 的 To 取自 Ship Info.ArrivalNotice；共用 CcTo 取自該通知設定。修改共用 MailTo 不會替代每筆 Ship Info 的到貨通知名單。','其他種類依各自 MailTo／CcTo 來源；保留其他通知名單與 SMTP、FromAddress、Enabled、Frequency、SendTimes 等未申請項目。','保存設定差異並核對有效來源；既有排程可能保存快照，不假設 Pending 或歷史紀錄全部被更新。設定消費與重啟影響見 [[20.4]]。','本地以 NotificationSmtpChecks 核對 SMTP 與 fallback，使用隔離排程檢查，不寄出真實郵件。使用者需變更名單時交由管理者更新，唯讀通知頁不能直接編輯。']:b(t)
h('12.2 日期與避免重複')
p('Delay Notification Date 是通知日期；郵件 Delivery Date 取自 Delivery Information 的 Delivery Date。Arrival Notice 的預約時間依台北時區設定的 SendTimes，非任意輸入排程時間。管制品邏輯另看 Flag、EL_FLAG 與排程服務，不以 ETA 等同寄送日。')
p('Pending／Sending／Sent／Failed 有不同含義；已寄信不能透過修改 DB 狀態撤回。SMTP 已接受而後續回寫中斷時，人工重送有重複風險，應先查 Log、排程與信件系統紀錄。')

chapter(13)
h('13.1 站台與執行身分')
p('ICP、Hub、FileGenerator 分別發布並核對實體目錄、Binding、App Pool、Runtime／Hosting Bundle。服務帳號需有資料庫及相應目錄權限；UNC 同時核對 Share 與 NTFS。不要使用互動登入者的 Z: 等映射磁碟作為長期服務路徑。')
h('13.2 ICPFileGenerator 需要呼叫 status 才開始的排查')
table(['檢查','判讀與處理'],[
 ('App Pool 是否 OnDemand／閒置停止','按 status 後才工作可能只是處理序被喚醒，不足以判定死鎖。對照程序啟動時間、IIS recycle 與 Worker Started Log。'),
 ('Application Initialization','確認 Windows Server IIS 應用程式開發功能已安裝 Application Initialization；依站台政策設 App Pool Start Mode=AlwaysRunning、應用程式 Preload Enabled=true。'),
 ('Idle／回收','需要持續背景工作時評估 Idle Time-out=0；排定回收仍會重啟，需觀察重啟後自動恢復。不要把關閉所有回收當成唯一修復。'),
 ('啟動是否真的成功','核對 Runtime、必要 TEL JSON、DB／路徑、事件檢視器與應用程式 Log；預載無法修復啟動例外。'),
 ('驗證方式','受控回收後不手動呼叫 status，觀察 Worker 啟動及既有／核准測試 Job 推進，再查 API。API 可回應與 Worker 進度要分開驗證。')
],[2.0,5.0])
p('IIS 維運建議參考 Microsoft 官方文件的 Application Initialization Module and Idle Timeout。若未來改為 Windows Service，需另外加入相應 Hosting 支援、安裝與恢復策略；目前專案不能只靠文件就視為已支援或已完成服務化。')
p('官方參考：https://learn.microsoft.com/en-us/aspnet/core/host-and-deploy/iis/?view=aspnetcore-9.0')

chapter(14)
h('14.1 建置與發布產物')
code(['dotnet publish ICP/ICP.csproj -c Release -o <STAGING>/ICP',
      'dotnet publish TELIntegrationHub/TEL.IntegrationHub.csproj -c Release -o <STAGING>/Hub',
      'dotnet publish ICPFileGenerator/ICPFileGenerator.csproj -c Release -o <STAGING>/FileGenerator'])
p('上列是待執行範例，<STAGING> 換成獨立暫存目錄。不要以複製幾個 .cs 原始檔取代網站發布；主機若用 VS 啟動原始碼，仍需重新建置並確認啟動的是更新後專案。')
table(['產物','核對內容'],[
 ('網站發布目錄','DLL、deps/runtimeconfig、web.config、static web assets、wwwroot、Config、Files、Resources。相依套件及 Razor 頁面依 publish 輸出。'),
 ('環境設定','部署該主機實際載入的 TEL JSON；保護帳密及原有通知名單。example 不會自動取代實際設定。'),
 ('DB／檔案','先確認相容性，備份需要的 DB／資料／設定；保存 uploads、Logs、輸入輸出，不隨站台覆蓋清空。'),
 ('文件','ICP-User-Guide.docx 與 ICP-Administrator-Guide.docx；CustomLinks 含正確環境／虛擬路徑及版本參數。')
],[1.5,5.5])
h('14.2 發布順序與驗證')
for t in ['記錄本次 commit／發布編號、三服務版本、SQL 與設定差異及回復方案。','備份並執行本版必要的相容性 SQL；移除舊表需另依相依順序，不能一律套同一順序。','於維護窗口停止或協調背景工作，部署完整發布產物，套用受保護環境設定與資料夾權限。','啟動依賴服務及三應用程式；核對 Log 的機器、AppSettings 與 DB／共用路徑。','驗證登入、權限、查詢、Settings、匯入、報表、附件與文件下載；經核准再測試跨服務起案、產檔及寄信。','使用 RequestId／messageId 追蹤到最後狀態；發布紀錄附驗證證據，不能只記首頁可開。']:b(t)
h('14.3 回復')
p('回復先停止相關寫入與 Worker，確認新資料是否需保留。還原相容的程式及設定；DB schema／資料依備份與核准腳本處理。回復後檢查未完成 Job、Outbox、寄送中的通知，避免舊新版同時處理或重複副作用。')

chapter(15)
table(['問題入口','完整處理位置'],[
 ('舊畫面／欄位不符','先核對站台與發布來源；[[14.1]]、[[26.2]]。'),
 ('例外／瀏覽器錯誤','依實際訊息搜尋 [[15.2]]，附 [[15.1]] 的診斷證據。'),
 ('起案／產檔／通知流程卡住','分別使用 [[21.1]]、[[21.2]]、[[21.3]]；恢復前查 [[26.4]]。')
],[2.1,4.9])
h('15.1 診斷證據包')
p('收集發生時間與時區、環境網址、電腦／處理序、發布編號、AppSettings 檔名、功能、Invoice／RequestId／messageId、錯誤堆疊及前後 Log。Network 另附 Status、Initiator、Response；移除密碼、Cookie、完整認證與不必要個資。')
p('ICP Program.cs 額外寫入 ContentRoot/Logs/icp-.log，保留 31 個檔案；Hub 與 FileGenerator 的 sink 依各自程式及環境設定核對。不要假設所有服務都寫相同檔名或同一目錄。')
h('15.2 實際錯誤訊息與修正方向').paragraph_format.page_break_before=True
p('以下字串來自既有問題紀錄或程式檢查；為方便搜尋保留英文原文，變動的號碼、列號與名稱使用占位符。不是本次重新造成或寫入主機的錯誤。',kind='維運建議')
table(['搜尋的實際訊息或片段','原因判斷與修正方式'],[
 ('GET …/lib/datatables/images/sort_both.png 404 (Not Found)','Network 查 Initiator 及 DataTables CSS 版本，修正套件資產路徑或統一樣式；不可用空圖片掩蓋。'),
 ('Permissions policy violation: unload is not allowed in this document.','常見來源為 VS Browser Link；分開檢查業務 API 及 JS 初始化，不能據此判斷資料庫故障。'),
 ('System.Threading.Tasks.TaskCanceledException: A task was canceled.','查取消 Token、停止除錯／回收、請求取消及 SQL／網路同時段 Log。正常停止與故障要分辨，不一律吞掉例外。'),
 ("Invalid column name 'WBS_ELEMENT'／無效的資料行名稱 'WBS_ELEMENT'",'核對實際 DB 及 017 腳本；ALTER 後同批 UPDATE 的編譯問題用既有動態 SQL 版本處理，見 [[7.1]]。'),
 ("System.IO.IOException … bootstrap.rtl.css.map … 另一個處理序正在使用檔案",'定位鎖定來源與重複建置／同步；協調解除占用後再建置，避免全域終止程序或刪除來源目錄。'),
 ('第 <列號> 列 DELIVERY_TO … Settings …','現行 Add DI/SA 依 DeliveryToList 的 Customer Code 比對；查該頁建檔值與欄位，不要以 Name 或 Key 替代；規則見 [[8]]。'),
 ('Required Shipping advice column …','檢查第一個分頁的表頭、版本及必要欄；修正輸入檔再測，不靠後面同名分頁繞過檢查。'),
 ('Please check County of origin','保留程式既有拼字以便搜尋；檢查來源國別與設定對照，修正後再產檔。見 [[11]]。'),
 ('File not found','先分清附件、歷史原檔或產出檔，再查實體路徑、檔案存在、服務身分及共用目錄權限；不要直接重匯全部資料。'),
 ('POST /api/messages/{messageId}/retry → 501','此入口尚未實作，不是補檔即可恢復；依 [[25.1]] 與 [[26.4]] 查專用流程。')
],[3.0,4.0])


log_examples(table,p,h,code)

chapter(16)
p('測試程式位於 tests，通常為可執行 checks 專案。本版本地端測試結果與紀錄見第 25 章；以下提供日後執行範例，未列入結果的專案不視為已通過。先讀 Program.cs 與相依項目確認是否操作外部服務，再在隔離環境執行。')
code(['dotnet run --project tests/AddDiSaImportChecks/AddDiSaImportChecks.csproj',
      'dotnet run --project tests/ForwarderUploadChecks/ForwarderUploadChecks.csproj',
      'dotnet run --project tests/SettingsExcelChecks/SettingsExcelChecks.csproj',
      'dotnet run -c Release --project tests/FileGenerationChecks -- "<20260916 sample.xlsx>" "<temporary output>"',
      'dotnet run --project tests/NotificationSmtpChecks/NotificationSmtpChecks.csproj'])
table(['變更類型','至少驗證'],[
 ('匯入','正確案例、缺必填、錯型別、重複鍵、Settings 比對、同 Invoice 一致性、失敗保留原資料。'),
 ('報表／權限','多值與空白、日期邊界、分頁／匯出一致、不同角色及 DataScope、資料列少時 Filter 不裁切。'),
 ('起案／產檔','無效快照阻擋、識別串聯、重試與逾時、重複副作用、下載內容／PDF 版面。'),
 ('通知','From fallback、各類 To／Cc 來源、時區／排程／Pending 取消、SMTP 失敗及寄送紀錄。'),
 ('部署','publish 產物完整、正確 TEL JSON、schema、IIS 自動啟動、共用檔與 Log。')
],[1.6,5.4])
h('16.1 現有檢查專案索引')
for group in [sorted(p.parent.name for p in (ROOT/'tests').glob('*/*.csproj'))[i:i+6] for i in range(0,len(list((ROOT/'tests').glob('*/*.csproj'))),6)]:
    p('、'.join(group))
p('測試 checks 原始碼供 CI 或開發驗證，不需複製進網站目錄。測試通過的數字應以當次輸出記錄，勿沿用舊次數當新版驗證。')

chapter(17)
p('新增欄位、Filter、Excel 卡控、通知及整合功能的同步範圍與本地測試集中於 [[17.2]]；本章不重複列維護相同清單。',kind='維運建議')
h('17.1 交付紀錄')
b('每次需求記錄問題、修改後行為、變更檔案與資料庫影響、驗證、部署順序及回復限制。資料／程式／設定／範本要分別列出。')
b('文件有兩種讀者：操作手冊教使用者如何完成及求助；本管理者手冊教開發與維運如何定位、修改與恢復。UI、卡控、排程與部署異動時同步更新相關章節。')
b('文件連結需有版本參數以降低舊快取；新文件改名同時更新網站實體檔與 Sidebar.CustomLinks。依主機虛擬路徑確認 /ICP，不能讓正式網站連到 localhost。')

change_matrix(table,p,h)

chapter(18)
h('18.1 主要程式來源')
sources=[('ICP/Program.cs','設定選用、DI、Workers、schema 確保與 middleware'),
 ('ICP/Data/ApplicationDbContext.cs','業務／權限資料表對映'),
 ('ICP/Services/PageDataScopeService.cs','資料範圍及無效規則行為'),
 ('ICP/Controllers/Setting/SystemConfigControllerBase.cs','分頁 Excel 匯入交易'),
 ('ICP/Services/NotificationMail/NotificationMailSender.cs','寄送紀錄與狀態'),
 ('ICP/Models/NotificationMail/NotificationMailOptions.cs','SMTP 與通知設定契約'),
 ('TELIntegrationHub/Program.cs','Hub 設定選用與啟動'),
 ('TELIntegrationHub/Controllers/ExportApiController.cs','Export 接收與回寫 API'),
 ('TELIntegrationHub/Services/ExportOrchestrationService.cs','建立 Job 與更新 Request'),
 ('TELIntegrationHub/Data/HubDbContext.cs','MESSAGE_LOG 對映；已無 RoutingRule'),
 ('ICPFileGenerator/Program.cs','FileGenerator DI／Worker 註冊'),
 ('ICPFileGenerator/Workers/FileGenerationWorker.cs','Claim、逾時、產檔與 Hub 回報'),
 ('ICPFileGenerator/Controllers/SystemController.cs','固定 status 的診斷限制')]
assert all((ROOT/path).exists() for path,_ in sources)
table(['來源','核對內容'],sources,[4.5,2.5])
h('18.2 文件版本')
table(['版本','日期','變更'],[
 ('1.0','2026-09-12','原 ICP 系統設定管理手冊，以 appsettings 及設定部署為主。'),
 ('2.0','2026-10-26','改名管理者手冊；擴充三服務架構、開發、部署與維運。')
 ,('2.1','2026-10-27','區分程式行為、維運建議與尚未實作；補流程及 WBS 實例。')
 ,('2.2','2026-10-27','採本地端測試；補部署、背景排查、恢復、備份與權限說明。')
 ,('2.3','2026-10-27','補快速索引、設定影響、錯誤及修改範圍；合併重複說明。')
 ,('2.4','2026-10-27','補關鍵字、Log、管理實拍與兩份範本；完成架構圖解及本輪優化。')
],[.65,1.2,5.15])
maintenance_record(table,p,h)
doc_update_checklist(table,p,h)
append_playbook(doc,table,p,b,h,code,ROOT)

# Navigable contents, return links and page numbers.
headings=[x for x in doc.paragraphs if x.style.name.startswith('Heading')]
targets={}
for i,para in enumerate(headings,1):
    a='admin_section_'+str(i);bookmark(para,a,i);targets[para.text]=a
import re
numbered={re.match(r'^(\d+(?:\.\d+)?)\s',title).group(1):anchor for title,anchor in targets.items() if re.match(r'^(\d+(?:\.\d+)?)\s',title)}
reference_count=0
paragraphs=list(doc.paragraphs)
for tab in doc.tables:
    for row in tab.rows:
        for cell in row.cells:paragraphs.extend(cell.paragraphs)
for para in paragraphs:
    if '[[' not in para.text:continue
    parts=re.split(r'\[\[(\d+(?:\.\d+)?)\]\]',para.text)
    assert len(parts)>1,para.text
    para.clear()
    for idx,part in enumerate(parts):
        if idx%2:
            assert part in numbered,part
            link(para,'第 '+part+' 節',numbered[part]);reference_count+=1
        elif part:para.add_run(part)
    if para._p.getparent().tag==qn('w:tc'):
        for run in para.runs:run.font.size=Pt(9.5)
assert reference_count>=60,reference_count
print('Cross references:',reference_count)
toc=targets['目錄']; active=False
for para in doc.paragraphs:
    if para.text=='目錄':active=True;continue
    if active and para.style.name.startswith('Heading'):break
    if active and para.text in chapters:
        text=para.text;para.clear()
        para.paragraph_format.tab_stops.add_tab_stop(Pt(485),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link(para,text,targets[text]);para.add_run('\t');field(para,'PAGEREF '+targets[text]+' \\h')
from docx.text.paragraph import Paragraph
for para in headings:
    if para.text=='目錄':continue
    nav=OxmlElement('w:p');para._p.addnext(nav);np=Paragraph(nav,para._parent)
    np.paragraph_format.keep_with_next=True;np.paragraph_format.space_after=Pt(3)
    link(np,'回目錄',toc)
    if para.style.name=='Heading 1':
        np.add_run('  ｜  ');link(np,'問題排查',targets[chapters[14]])
        np.add_run('  ｜  ');link(np,'部署流程',targets[chapters[13]])
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('ICP系統管理者手冊 2.4  ｜  第 ');field(footer,'PAGE');footer.add_run(' 頁／共 ');field(footer,'NUMPAGES');footer.add_run(' 頁')
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
doc.core_properties.title='ICP系統管理者手冊'
doc.core_properties.subject='開發架構 程式維護 設定 部署與異常排查'
doc.core_properties.author='ICP Project Team'
OUT.mkdir(exist_ok=True,parents=True);doc.save(DOCX);make_templates(OUT);print(DOCX)
