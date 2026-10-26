"""Evidence-labelled developer and operations supplements for admin guide 2.4."""
from PIL import Image, ImageDraw, ImageFont
from docx.shared import Inches
from icp_admin_findability import setting_impact
from icp_admin_completion import architecture, screenshot_pages

EXTRA_CHAPTERS=['19 架構流程圖','20 設定與生效驗證對照','21 故障排查作業流程',
                '22 發布與回復案例','23 新增 Detail 欄位實例','24 資料關聯與狀態轉換','25 開發接手與功能邊界','26 維運實務與閱讀路徑']

def add_reading_key(doc,table):
    doc.add_heading('1.2 三種內容標示',level=2)
    table(['標示','讀法','不能推論的事情'],[
        ('目前程式行為','已由本次工作區的程式、Options、SQL 或測試來源確認；以來源檔為依據。','本版採本地端測試；不把隔離測試描述成真實寄信或下游寫入。'),
        ('維運建議','建議人員執行的檢查、變更、驗證及回復程序。','不是程式自動完成的功能；範例 SQL、指令與測試資料須先核對環境。'),
        ('尚未實作','目前指定入口未提供或只是占位行為，須另外開發與驗證。','不表示所有相近功能都未實作，也不是已承諾的交付功能。')
    ],[1.2,2.9,2.9])
    doc.add_paragraph('本版移除正式環境驗收要求，以本地端測試結果為完成依據。「尚未實作」仍保留為功能邊界；表內操作與核對欄為維運建議。')

def flow_image(path,title,nodes):
    width=1500; cols=3; rows=(len(nodes)+2)//3; height=100+rows*185
    im=Image.new('RGB',(width,height),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/msjh.ttc',28)
    bold=ImageFont.truetype('C:/Windows/Fonts/msjhbd.ttc',32)
    d.text((25,18),title,font=bold,fill='#183c59')
    boxes=[]
    for i,text in enumerate(nodes):
        row,col=divmod(i,cols); col=col if row%2==0 else 2-col
        x=30+col*490;y=90+row*185;box=(x,y,x+425,y+130);boxes.append(box)
        d.rounded_rectangle(box,radius=14,fill='#edf4fa',outline='#1f4e78',width=3)
        lines=[]
        for line in text.split('\n'):
            current=''
            for c in line:
                if d.textlength(current+c,font=font)>385:lines.append(current);current=c
                else:current+=c
            lines.append(current)
        for j,line in enumerate(lines):d.text((x+15,y+12+j*34),line,font=font,fill='#172b39')
        d.ellipse((x-10,y-16,x+34,y+28),fill='#1f4e78');d.text((x+3,y-10),str(i+1),font=font,fill='white')
    for i in range(len(boxes)-1):
        a=boxes[i];b=boxes[i+1]
        if a[1]==b[1]:
            right=b[0]>a[0];start=(a[2] if right else a[0],a[1]+65);end=(b[0] if right else b[2],b[1]+65)
            d.line((start,end),fill='#1f4e78',width=4);x,y=end;sign=-1 if right else 1
            d.polygon([(x,y),(x+sign*15,y-9),(x+sign*15,y+9)],fill='#1f4e78')
        else:
            x=(a[0]+a[2])//2;d.line((x,a[3],x,b[1]),fill='#1f4e78',width=4)
            d.polygon([(x,b[1]),(x-9,b[1]-15),(x+9,b[1]-15)],fill='#1f4e78')
    im.save(path)

def append_playbook(doc,table,p,b,h,code,root):
    assets=root/'outputs/manual/admin-v21-diagrams';assets.mkdir(exist_ok=True)
    def chapter(index):doc.add_heading(EXTRA_CHAPTERS[index],level=1).paragraph_format.page_break_before=True
    def note(kind,text):doc.add_paragraph('【'+kind+'】'+text)
    def picture(name,title,nodes):
        path=assets/name;flow_image(path,title,nodes)
        doc.add_paragraph().add_run().add_picture(str(path),width=Inches(6.7))
    architecture(doc,table,p,h,root,flow_image)

    chapter(1)
    note('目前程式行為','下列預設值來自 Options 類別，不是主機目前的有效值。必要性分為啟動必要、功能使用必要及可選；環境 JSON、環境變數與命令列仍可能覆寫。秘密值只說用途，不列實際內容。')
    groups=[('20.1 ICP 設定',[
        ('App.Mode／SuperUser','可選；PRD／Off','啟動身分與權限設定；變更後重新啟動／登入核對。'),
        ('ConnectionStrings:ICP_Connection／ILC_Connection／FIESTA_Connection','功能相依；無可用正式預設','分別驗證業務查詢、登入組織與郵件群組；不得只看首頁。'),
        ('ForwarderDataUpload:StoragePath／MaxSizeMb','功能必要；uploads/forwarder／50','重啟後核對服務身分可讀寫；測試原檔保存與下載。'),
        ('TariffData:StoragePath／MaxSizeMb','功能必要；uploads/broker/tariff／50','核對三類附件路徑與檔案大小限制。'),
        ('Integration:RabbitMq:Enabled','可選；false','Outbox Worker 監看 Options；false 保留待發，不是已完成。變更連線仍建議重啟。'),
        ('RabbitMq:HostName／Port／VirtualHost／Exchange／RoutingKey','功能必要；localhost／5672／/／tel.integration／icp.shipinfo.case.initiated','密碼由環境供應；核對 Broker、綁定及一筆核准測試事件。'),
        ('Integration:Outbox:PollIntervalSeconds／MaxRetryCount／BatchSize','可選；10／5／20','Worker 讀 IOptionsMonitor；核對新批次間隔與重試行為。'),
        ('Integration:Hub:BaseUrl','Export 必要；http://localhost:5261','依 HTTP Client 註冊及環境核對，正式變更後重啟驗證。'),
        ('Integration:Export:InputDirectory／OutputDirectory','Export 必要；ICPFileGenerator\\Input／ICPFileGenerator\\Output','解析 ContentRoot 後核對共享實體路徑與服務身分。'),
        ('Sidebar:CustomLinks','可選；依 JSON','Name、Url、Enabled、OpenInNewTab、SortOrder；重開頁面並實際開檔核對版本。'),
        ('LocalizationManagement:Enabled','可選；false','核對功能可見性與 Resources 寫入權限；不是僅改顯示字。')]),
        ('20.2 Hub 與 FileGenerator',[
        ('Hub ConnectionStrings:HubDatabase','功能必要；程式有 LocalDB fallback','維運建議：正式環境必須明確設定；查 MESSAGE_LOG 所在庫。'),
        ('Hub ICP_Connection／FIESTA_Connection','啟動必要；無有效值會拒絕','啟動 Log 與依賴服務註冊；核對 ICP 完成回寫及 ARUR 操作者查詢。'),
        ('Hub ILC_Connection／ICPFileGenerator','對應功能必要','起案寫入與 Export Job 使用；同名庫不是同一實體庫的保證。'),
        ('Hub Integration:RabbitMq:Enabled／QueueName','可選；false／tel.integration.hub.queue','Enabled 在啟動時決定是否註冊 Bus，需重啟；查 queue 與 consumer。'),
        ('Hub RabbitMq:ExchangeType／RoutingKey(s)','topic；Options 有 RoutingKey(s)','目前註冊綁定使用事件常數；僅改 RoutingKeys 不代表實際 Queue 綁定已改。'),
        ('Hub Integration:Database:EnsureCreatedOnStartup','可選；true','啟動時讀取；EnsureCreated 不是逐版 migration 或所有既有欄位升級。'),
        ('FileGenerator ConnectionStrings:ICPFileGenerator','功能必要；無有效連線不能 Claim','查 Job 與 Request 是否同環境；不要只測 API status。'),
        ('FileGenerator:PollingIntervalSeconds／WorkerId','可選；10／ICPFileGenerator-01','Worker 使用啟動 IOptions 快照；變更後重啟，查 Started／Claimed Log。'),
        ('FileGenerator:OutputDirectory','功能必要；..\\ICP\\ICPFileGenerator','預設相對位置不保證等同 ICP OutputDirectory；核對完整路徑。'),
        ('FileGenerator:ProcessingTimeoutMinutes／MaxRetryCount','可選；30／3','重啟後核對逾時工作處理；不保證所有 Failed 自動回復。'),
        ('FileGenerator:Hub:BaseUrl','回報必要；http://localhost:5261','重啟後核對回報結果及 EXPORT_REQUEST 更新。')]),
        ('20.3 通知信',[
        ('NotificationMail:Smtp:Host／Port／EnableSsl','smtp.tel.co.jp／25／false','目前核准 Relay 參數；查連線與寄送紀錄，不會自動測試投遞。'),
        ('Smtp:FromAddress／FromName','空寄件者 fallback；名稱可空','FromAddress 空白、null、未設時為 TETIS87GEMSAdmin@tel.com；既有非空值保留。'),
        ('Smtp:UserName／Password','可空；預設空','只有 UserName 非空才設憑證；不能因為有 Password 就推論已使用認證。'),
        ('各通知 Enabled／Frequency／SendTimes','false／Daily／空清單','依通知類別的實作判斷觸發／排程，不能只靠 Options 註解假設所有 Frequency 都支援。'),
        ('各通知 Title／MailTo／CcTo','主旨空／名單空','啟用通知需有有效目的名單；ArrivalNotice To 取 Header，其共用 CcTo 取設定。'),
        ('NotificationMail:ShipInfoUrl','可選；空','核對信內連結指向正確環境，含 IIS 子應用程式路徑。')])]
    for title,rows in groups:
        h(title);table(['目前程式行為：設定鍵','必要性與程式預設','維運建議：生效後驗證'],rows,[2.7,1.7,2.6])
    note('維運建議','正式設定變更統一記錄：環境／實際檔名、舊值摘要、新值摘要、Secret 來源、生效方式與驗證結果。不要在記錄中貼密碼或完整連線字串。跨服務連線與路徑同時比對。')

    setting_impact(table,p,h)

    chapter(2)
    note('維運建議','以下是人員排查順序，並非程式自動修復流程。先用唯讀資料定位；未確認下游是否完成前，不執行重送或狀態重設。')
    for title,rows in [
        ('21.1 起案卡住',[
          ('1 定位請求','收集 Header／Invoice、CaseType、messageId、時間。查 Outbox 狀態及 LastError。','沒有 Outbox：回到起案驗證／提交；有則繼續。'),
          ('2 Pending／Failed','查 ICP Worker、RabbitMq.Enabled、Broker／路由、RetryCount。','發布錯誤先排除，勿新建同一業務案件；Failed 的回排入口另查原程式。'),
          ('3 Published','查 Hub MESSAGE_LOG 同 messageId、Consumer／queue、ILC 目的資料。','Published 只代表發布；下游已寫入而未回寫時避免重送。'),
          ('4 完成回寫','查 Hub 的 ICP_Connection、Outbox 狀態與案件狀態。','MarkCompleted 預期 Published；不符時可能略過，依 Log 排查而非直接改 Completed。')]),
        ('21.2 產檔不動',[
          ('1 Request','查 RequestId、EXPORT_REQUEST 的 Status／ErrorMessage。','Hub 未收到或找不到 Request 時，先核對庫及 BaseUrl。'),
          ('2 Job','用 request/{requestId} 查 Job；查 Worker Started／Claimed。','無 Job 查 Hub 接受流程；Pending 不動查程序、Claim SQL 及 DB。'),
          ('3 Processing','核對 StartTime、逾時參數、輸入路徑、讀寫權限與產檔錯誤。','重啟或回排前先確認沒有仍在工作的產檔程序。'),
          ('4 Completed／Failed','比對 Job、輸出檔、Hub 回報與 Request 狀態。','Job 完成而 Request 未完成：查 callback，不是先重新產檔。')]),
        ('21.3 通知未寄',[
          ('1 觸發資料','確認通知種類、Enabled、業務 Flag、原因、Header 地址及日期。','條件未滿足時不是 SMTP 故障。'),
          ('2 排程','查 Pending／Sending／Sent／Failed 與台北時段；查 Worker。','Pending 可依流程取消；Sending 不直接重建。'),
          ('3 寄送紀錄','查 NOTIFICATION_MAIL_LOG 的 MailType、EventKey、時間及錯誤。','沒有 attempt 查前段；Failed 查 SMTP；Sent 再查信箱投遞。'),
          ('4 恢復','收件者／設定修正後檢查既有排程快照；保留原 Log。','SMTP 接受後回寫失敗可能產生重複風險，人工核對後才決定重送。')])]:
        h(title);table(['維運建議：步驟','查詢位置與證據','恢復條件'],rows,[1.35,2.9,2.75])
    note('尚未實作','Hub 的 POST /api/messages/{messageId}/retry 回 501；FileGenerator status 沒有 DB／Worker 心跳檢查。不要把這兩個入口當成自動回復與完整健康檢查。')

    chapter(3)
    note('維運建議','以下以現有 WBS_ELEMENT 改修為教學發布案例，不代表本次已再次執行 SQL 或發布正式主機。這是一個向後相容欄位增加案例，不能照搬到刪欄等破壞性變更。')
    table(['階段','應執行與核對','成功證據／停止條件'],[
        ('版本與備份','記錄 commit／發布編號；備份 ICP DB、目前發布目錄、受保護設定及受影響資料；確認回復窗口。','備份可用且可辨識同一環境；未備妥不執行 schema 異動。'),
        ('SQL','核對 017_AddDetailWbsElement.sql 的目標庫；先新增 Detail 欄，再以動態 SQL 回填。','欄位長度 30，既有非空 Detail 不被覆寫；檢查回填筆數與樣本。'),
        ('程式與 UI','發布 Entity、DTO、Catalog、Mapper、匯入規則、Config、Resources 及相關功能的完整產物。','不是只複製 IcpDetail.cs；新舊讀取與明細保存結果正確。'),
        ('設定與範本','本案例無須改 SMTP／MailTo／CcTo；新版 DI/SA 範本含 WBS_ELEMENT。','實際下載表頭與程式一致；環境 Secret、既有名單保持。'),
        ('驗收','用隔離測試驗證空值、30／31 字、兩列不同 WBS、匯入／保存／再查／匯出。','通過紀錄附版本；若只有預覽未 Save，不算保存驗收。'),
        ('回復','停相關寫入，回到相容舊版程式與設定；先保留新增欄及已寫入資料。','舊版仍讀 Header 或忽略 Detail 的影響須明示；不直接 DROP 欄位以免遺失新值。')
    ],[1.1,3.1,2.8])
    h('22.1 唯讀核對範例')
    note('維運建議','以下只查欄位與回填覆蓋情況；請在正確 ICP DB 執行。查到 NULL 是資料狀態，不一定是發布失敗；應比對原始 Header 與業務規則。')
    code(["SELECT COL_LENGTH(N'dbo.ICP_DETAIL', N'WBS_ELEMENT') AS Bytes;",
          'SELECT COUNT(*) AS DetailRows,',
          ' SUM(CASE WHEN WBS_ELEMENT IS NOT NULL THEN 1 ELSE 0 END) AS HasWbs',
          'FROM dbo.ICP_DETAIL;'])
    note('目前程式行為','nvarchar(30) 的 COL_LENGTH 是最大位元組長度（通常 60），不是字元數 30；不要把此查詢結果誤判為型別不符。')

    chapter(4)
    note('目前程式行為','現有 WBS_ELEMENT 是 Detail 可分列不同的選填欄位，最多 30 字；下列用已存在實作說明修改順序，不宣稱新增另一個未實作欄位。')
    table(['層次','現有檔案與核對內容','維運建議：增加欄位時'],[
        ('DB／Entity','sql/017_AddDetailWbsElement.sql、Models/Icp/IcpDetail.cs、ApplicationDbContext','先決定 Header／Detail 歸屬、型別、NULL、長度與回填，不只加畫面。'),
        ('資料傳遞','ShipInfoDetailRowDto、ShipInfoRowDtoMapper、ShipInfoFieldCatalog','新增映射並測試模型到 JSON 的值，避免表頭有欄而內容空白。'),
        ('上傳','AddDiSaExcelColumnMap、AddDiSaImportRules、AddDiSaImportService','核對必填、同 Invoice 差異白名單、大小寫／日期正規化、交易及 Save 再驗證。'),
        ('編輯與清單','shipinfo-detail-form-fields.json、shipinfo-pro-datatable-fields.json','確認欄位可編輯與保存、查詢可取回；Config 不能自行產生 DB 欄。'),
        ('匯出與郵件','相關資料查詢、匯出、ArrivalNoticeMailContent、Resources','依需求核對各輸出是否應包含新欄；不是所有下游都自動納入。'),
        ('測試與文件','AddDiSaImportChecks 等；操作手冊與部署紀錄','新增正常／邊界／拒絕案例，發版附 SQL、設定、範本及回復影響。')
    ],[1.2,3.0,2.8])
    note('維運建議','示例驗收矩陣：空白可接受、30 字可接受、31 字拒絕；同 Invoice 兩明細 WBS 不同可接受；非法必填資料仍整檔拒絕；保存後再查與匯出不遺失值。新增別的欄位時，應依新規格重定義矩陣。')

    chapter(5)
    note('目前程式行為','以下區分程式關聯及來源已定義的約束。正式資料庫是否已套用約束仍須查 sys.foreign_keys／sys.indexes；不能只憑模型認定主機一致。')
    table(['關聯／索引','實際性質','來源與限制'],[
        ('ICP_HEADER ↔ ICP_DETAIL','多處以 InvoiceNo＋TetPo 關聯；目前 EF 未宣告此關聯外鍵','ApplicationDbContext、PageDataScopeService、017 回填。不同流程分組鍵不能一律換成這一組。'),
        ('EXPORT_REQUEST ↔ ICPFileGeneratorJob','透過 RequestId 程式關聯；Job SQL 有 RequestId UNIQUE','FileGenerator sql/001 未宣告到 EXPORT_REQUEST 的 FK；表存在不保證已執行同版 SQL。'),
        ('Outbox ↔ Hub MESSAGE_LOG','messageId 相關聯，屬跨服務查詢識別','兩個 DB 間沒有這種外鍵保證；不能以一邊存在推論另一邊完成。'),
        ('排程 ↔ Header','HeaderId 程式關聯；Pending HeaderId 有 filtered unique index','Arrival／Controlled Schedule 的 Pending 唯一性不是所有歷史列都唯一。'),
        ('角色授權','RolePermissions → Roles／Resources 等 EF HasForeignKey','角色指派另有唯一組合索引；實際 DB 需核對 schema。')
    ],[2.0,2.5,2.5])
    h('24.1 狀態轉換').paragraph_format.page_break_before=True
    table(['對象','目前程式行為','維運建議與限制'],[
        ('Outbox','Pending → Published；發布失敗增加 RetryCount，未達上限保留 Pending，達上限 Failed；Hub 完成回寫預期 Published → Completed。','Failed 原程式有 TryRequeueFailedAsync 回 Pending；與 Hub 501 retry API 是不同機制，不能混用。'),
        ('Job','Pending → Claim Processing → Completed／Failed；Processing 逾時依次數回 Pending 或 Failed。','一般產檔 Failed 不代表自動重新產檔；查 JobRepository 與實際條件。'),
        ('Export Request','Hub 接受後 Processing；完成或失敗回報更新 Completed／Failed。','與 Job 狀態可能短暫不同；先查 callback，不直接重建 Request。'),
        ('Hub MESSAGE_LOG','Processing → Success 或 Failed；查詢 errors 另包含 DeadLetter。','查詢涵蓋 DeadLetter 不代表每個 Consumer 都會寫入它。'),
        ('通知排程／紀錄','排程 Pending、Sending、Sent、Failed／Cancelled 依服務處理；通知 Log 先 Sending 再 Sent 或 Failed。','排程與寄送 Log 為兩種紀錄；Sent 不等於對方信箱已投遞成功。')
    ],[1.35,3.1,2.55])

    chapter(6)
    note('維運建議','接手開發者依下列順序進行本地端核對；本版不建立正式環境驗收表。使用者於 2026-10-27 告知 Hub 與 ICPFileGenerator 已啟動，本文件記錄此狀態，不要求重新啟動服務。')
    table(['檢查點','需填寫的證據','通過標準'],[
        ('環境','機器名、SDK／Runtime、啟動方式、實際 AppSettings 檔名','三服務框架符合；非 AGA-PC 會選 TEL 設定，沒有誤用正式連線。'),
        ('資料與檔案','各連線的 DataSource／Database 摘要、服務身分、完整輸入輸出與附件路徑','在核准測試環境可查資料及讀寫測試檔；無密碼留存。'),
        ('最小啟動','各程序 Started Log、Hub health、FileGenerator status 與 Job 查詢','API、DB、Worker 分別驗證；status 200 不是全部健康。'),
        ('單元與隔離測試','所改功能的 checks 專案、組態、版本、結果與 Log','測試程式外部副作用已先讀取；不以舊次數當新版結果。'),
        ('本地流程測試','測試模型、RequestId／messageId、模擬依賴與結果','核對成功、拒絕、失敗及重複情境；不以模擬結果宣稱真實下游寫入。'),
        ('交接','發布包、SQL 清單、設定差異、回復方式、操作與管理手冊版本','下一位人員能定位相同檔案、資料與狀態；聯絡窗口由公司提供。')
    ],[1.1,3.1,2.8])
    h('25.1 尚未實作清單').paragraph_format.page_break_before=True
    table(['分類','指定功能或入口','來源與替代處理'],[
        ('尚未實作','Hub 通用 retry API','MessagesController 明確回 501。人工排查及既有個別回排需依專用流程，不能呼叫這個 API 期待恢復。'),
        ('尚未實作','FileGenerator 完整依賴／Worker 心跳健康檢查','SystemController 僅固定回 Running／version，未驗證資料庫與 Worker。以唯讀 Job 與 Log 分段驗證。'),
        ('尚未實作','僅依文件即可安裝 FileGenerator Windows Service','目前未配置 AddWindowsService／UseWindowsService。服務化需另加 Hosting、安裝與恢復測試。'),
        ('目前程式行為','NoOpTargetWriter 為占位元件','Hub 雖註冊此通用介面，不代表既有 Deposit／ARUR 真實寫入未實作；目前有專用 ILC 寫入服務。'),
        ('維運建議','新功能／監控強化候選','完整 health、告警、通用安全重試可另開需求；本次未修改業務程式，不當作已交付功能。')
    ],[1.15,2.75,3.1])
    h('25.2 本地端測試紀錄').paragraph_format.page_break_before=True
    note('目前程式行為','沿用 2026-10-27 本地端測試紀錄；Release 組態，結果取自 outputs/manual/admin-v22-local-tests。本次文件優化另作內容、連結與版面檢查，不重複計入業務測試項數。測試使用隔離資料、暫存檔與模擬依賴；未以真實 SMTP 寄信，未提交正式起案或重送業務事件。')
    import re
    rows=[]
    for name in ['AddDiSaImportChecks','ForwarderUploadChecks','SettingsExcelChecks','TariffUploadChecks','TariffAttachmentChecks','ShippingReportChecks','MassDataReportChecks','NotificationSmtpChecks','FileGenerationChecks']:
        path=root/'outputs/manual/admin-v22-local-tests'/f'{name}.log'
        text=path.read_text(encoding='utf-8-sig',errors='replace') if path.exists() else ''
        # Completion and process exit are recorded separately by the test runner.
        status_path=path.with_suffix('.exit')
        assert status_path.exists() and status_path.read_text().strip()=='0',name
        lines=[line.strip() for line in text.splitlines() if re.search(r'passed|通過|PASS|checks',line,re.I)]
        result=lines[-1] if lines else '執行完成；Exit Code 0'
        rows.append((name,result))
    table(['本地端測試專案','實際輸出摘要'],rows,[3.15,3.85])
    note('維運建議','後續修改程式時，重跑受影響的測試並附當次 Log；文件或設定連結更新需另外核對版本、檔案雜湊、目錄連結與頁碼。既有較早核對紀錄保留供追溯，不作本版的正式環境待辦清單。')
    append_operations(doc,table,p,b,h,note,root)

def append_operations(doc,table,p,b,h,note,root):
    doc.add_heading(EXTRA_CHAPTERS[-1],level=1).paragraph_format.page_break_before=True
    h('26.1 閱讀路徑與流程維護位置')
    note('維運建議','依症狀及角色查閱 [[1.3]]；設定修改看 [[20.4]]，開發修改範圍看 [[17.2]]，每次維護記錄使用 [[18.3]]。以下保留部署分類與恢復界線，完整通知與 IIS 排查不重複刊載。')
    h('26.2 部署清單必須用資料夾與檔案')
    note('維運建議','每次改修建立一份清單，列實際相對路徑、用途、目標服務與順序。下表是清單格式，不是要把整個工作區或以下所有檔案重新上傳。')
    table(['資料夾或檔案類別','交付方式與注意事項'],[
        ('Controllers／Services／Models／Repositories','原始碼提交 CI/CD；主機使用完整 publish 產物，僅複製 .cs 不會更新執行中的程式。'),
        ('Views／wwwroot／Config','確認 csproj 的發布規則與 publish 實際內容；檢查 JS／CSS 版本及瀏覽器快取。'),
        ('Files/*.xlsx、wwwroot/docs/*.docx','範本與手冊分開列檔名；下載後比對表頭或 SHA-256。不要遺漏實際供下載的副本。'),
        ('appsettings*.json','只更新本次需要的區段；保留機器專用連線、路徑、MailTo／CcTo、啟用與排程。'),
        ('sql/*.sql','依版本逐支列出順序及執行一次或可重跑條件；先確認目標庫與相容性。'),
        ('tests/*、tools/*、outputs/*','供原始碼、CI 與文件追溯；不整批複製到網站。網站僅部署指定下載文件。')
    ],[2.5,4.5])
    note('維運建議','清單另附備份位置、發布編號與回復版本。部署後出現舊畫面，先核對實際站台實體路徑、publish 版本、Config、Views／編譯產物與瀏覽器載入檔案。')
    h('26.3 背景工作排查入口')
    note('維運建議','Hub 與 ICPFileGenerator 已啟動是使用者提供的狀態。呼叫 status 後才工作或回收問題使用 [[13.2]]；Request／Job 沒前進使用 [[21.2]]。完整判讀集中於這兩處，恢復前再看 [[26.4]]。')
    h('26.4 重送前必須確認的界線')
    table(['流程','恢復前核對','不能直接推論'],[
        ('起案／Outbox','查 messageId、下游是否已寫入、重複保護與現有回排機制。','Hub 通用 retry API 未實作；不得當作可用按鈕。'),
        ('Export／Job','查 RequestId、現有輸出與完成回寫；區分 Processing 逾時與一般 Failed。','一般 Failed 不表示一定自動重試；有檔案也不表示 callback 完成。'),
        ('通知信','查 EventKey、排程、Log、SMTP 接受紀錄及目前收件者。','Sent 不保證投遞；SMTP 接受後 DB 回寫失敗可能已寄出，不能盲目重送。'),
        ('整批 Settings／匯入','確認目前資料、舊批次原檔及整檔驗證結果。','匯入成功不等於歷史原檔仍存在；重新上傳可能再次取代資料。')
    ],[1.4,3.0,2.6])
    h('26.5 通知設定變更入口')
    note('維運建議','修改 MailTo／CcTo、保留其他通知設定及既有排程快照的完整實例集中於 [[12.1]]；生效方式與本地測試見 [[20.4]]。管理者依申請更新，使用者不在唯讀通知頁直接修改。')
    h('26.6 備份與回復演練')
    table(['備份對象','本地演練方式','保留與回復限制'],[
        ('發布程式與設定','在獨立本地目錄保存前版 publish 與設定差異；比對版本、載入來源及文件雜湊。','不要在一般工作單或文件保存秘密值；回復須使用相容版本。'),
        ('資料庫','僅使用隔離測試資料庫演練備份、還原及 schema 查詢；留存筆數與關鍵值比對。','此為日後維運演練建議，本次未執行實庫備份還原；程式回復不會撤回資料。'),
        ('附件與原檔／輸出檔','在本地暫存目錄複製測試檔，比對檔名、大小、SHA-256 與讀取權限。','DB 與檔案需對得上；新增欄位或檔案先保留，不為回復而直接刪除。'),
        ('未完成背景工作','用隔離模型測試 Pending／Processing／Failed 及重複識別的處理；執行前讀取測試副作用。','真實運行恢復前另查 Job、Outbox 與通知紀錄，避免同時處理造成重複。')
    ],[1.55,2.9,2.55])
    h('26.7 權限與限制的維護')
    note('目前程式行為','Settings Excel 下載依可查看、上傳依可新增該頁資料的權限；單頁取代不影響其他頁，驗證失敗保留原資料。本地以 SettingsExcelChecks 核對；角色、後端授權與 Session 的完整說明見 [[9.2]]。')
    screenshot_pages(doc,table,p,h,root)
