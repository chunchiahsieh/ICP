"""Final usability, architecture and maintenance aids for administrator guide 2.4."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def keywords(table,p,h):
    h('1.4 搜尋關鍵字對照').paragraph_format.page_break_before=True
    p('使用 Ctrl＋F 搜尋任一中文、英文欄位、資料表或程式名稱；同一列指向相同主題。日期、檔案路徑及畫面例值不作搜尋前提。',kind='維運建議')
    table(['中文與畫面名稱','英文及程式關鍵字','主題位置'],[
        ('新增 DI/SA／上傳卡控','Add DI/SA、AddDiSa、Upload Data、INVOICE_SEQ','[[8.1]]、[[17.2]]'),
        ('交貨對象／客戶代碼','Delivery To、DeliveryToList、Customer Code、DELIVERY_TO','[[8]]、[[15.2]]'),
        ('貨代上傳紀錄／原檔','Forwarder Data Upload、Upload History、ForwarderDataUploadArchive','[[8]]、[[26.4]]'),
        ('報關附件／提單','Customs Clearance Data、TariffData、Declaration PDF、Import Doc','[[8]]、[[15.2]]'),
        ('押金起案／ARUR 起案','Deposit、ARUR、INTEGRATION_EVENT_OUTBOX、messageId','[[19.2]]、[[21.1]]'),
        ('匯出／產檔／工作','Export、FileGenerationWorker、EXPORT_REQUEST、ICPFileGeneratorJob、RequestId','[[19.3]]、[[21.2]]'),
        ('通知日期／交貨日','Delay Notification Date、Delivery Date、DeliveryDelay','[[12.2]]、[[19.4]]'),
        ('收件者／副本／排程','MailTo、CcTo、SendTimes、EventKey、NOTIFICATION_MAIL_LOG','[[12.1]]、[[15.3]]'),
        ('角色權限／資料範圍','Role permissions、ResourceCode、ActionCode、DataScope、Session','[[9.2]]、[[26.8]]'),
        ('篩選／多行／空白','Filter、Checkbox、Text、DateRange、pro-datatable','[[6.1]]、[[17.2]]'),
        ('啟動／健康檢查／閒置','status、Running、Preload、AlwaysRunning、Idle Time-out','[[13.2]]、[[25.1]]'),
        ('部署／版本／手冊更新','publish、appsettings、Config、SHA-256、Sidebar.CustomLinks','[[18.3]]、[[18.4]]、[[26.2]]')
    ],[1.75,3.3,1.95])


def log_examples(table,p,h,code):
    h('15.3 Log 判讀與識別串聯').paragraph_format.page_break_before=True
    p('以下為依目前程式訊息樣板製作的去識別化教學示例，不是本次服務實際執行證據。<REQ-A>、<MSG-A>、<JOB-A>、<EVENT-A> 代表替換後識別；時間採台北時間示例，實查須先確認 Log 的時區。',kind='維運建議')
    code(['[INF] [ICPFileGenerator.Worker] Started WorkerId=LOCAL-DEMO PollingIntervalSeconds=10',
          '[INF] [ICPFileGenerator.Worker] Claimed JobId=<JOB-A> RequestId=<REQ-A>',
          '[INF] [ICPFileGenerator.Hub] Notified Hub Completed RequestId=<REQ-A>',
          '      OutputFilePath=<OUTPUT>/demo.zip'])
    p('第一行只證明 Worker 已啟動；第二行表示已 Claim 工作；第三行在 Hub HTTP 成功回應後才記錄，仍需用同一 RequestId 核對 Export Request、Job 與輸出檔。Claimed 不是 Completed，固定 status 回 Running 也不能替代這些證據。',kind='維運建議')
    table(['識別與查詢鏈','如何串聯','不要混用'],[
        ('Export：RequestId → JobId','由 EXPORT_REQUEST 的 Id 找 Job.RequestId；Worker Claimed 行會同列 JobId 與 RequestId。回報 Hub 再用 RequestId 查 Request。','一個是請求識別、一個是工作識別；不是同一欄或必然相同值。'),
        ('起案：Outbox Id → MessageId','先找 Outbox 事件，再對照 payload 的 MessageId 與 Hub MESSAGE_LOG；InvNo／CorrelationId 輔助定位。','Invoice 可有多種事件；只用 Invoice 不足以判斷某一次是否完成。'),
        ('通知：MailType＋EventKey','用種類與事件鍵核對 NOTIFICATION_MAIL_LOG，再找相應排程及時間。實際收件人與信件內容只在授權範圍檢查。','沒有通用保證讓 EventKey 等於 RequestId 或 messageId；通知 Log 不是 Job 表。')
    ],[1.7,3.0,2.3])
    p('來源：ICPFileGenerator/Workers/FileGenerationWorker.cs、Services/HubNotificationService.cs；Hub 的 Consumers 與 ICP NotificationMailSender。通知 Sender 主要寫資料庫寄送紀錄，不應虛構一條共用「Mail sent」文字 Log 作完成條件。',kind='目前程式行為')
    h('15.4 失敗與取消的判讀範例').paragraph_format.page_break_before=True
    code(['[WRN] Failed to publish integration event <MSG-A>. RetryCount=2, Permanent=False',
          '[ERR] Failed handling Deposit case event <MSG-A>',
          '[ERR] [ICPFileGenerator.Worker] Job failed JobId=<JOB-A>',
          '[WRN] [ICPFileGenerator.Hub] Notified Hub Failed RequestId=<REQ-A> Error=<reason>'])
    p('這四行分別代表不同服務的示例，不是一組必然依序發生的事件。第一行是發布失敗，第二行是 Hub 下游處理失敗，第三行是產檔工作例外，第四行代表失敗結果已成功回報 Hub。排查時先確認是哪個階段，再讀同時間的例外堆疊；不要看到 Failed 就重送整份資料。',kind='維運建議')
    table(['訊息或狀態','判讀與下一步'],[
        ('Permanent=False','本次未達 Outbox 發布重試上限；核對 RetryCount 與設定，不是宣告已恢復。'),
        ('Hub MESSAGE_LOG = Failed','依 messageId 查目的系統是否已部分寫入；人工恢復先看 [[26.4]]。'),
        ('Job failed／Notified Hub Failed','核對 Job.ErrorMessage 與 Request；失敗回報成功不代表產檔成功。'),
        ('TaskCanceledException／OperationCanceledException','先核對停止除錯、IIS 回收、stoppingToken、請求取消與 DB timeout。FileGenerationWorker 外層在服務停止時可正常離開；工作中取消仍要查 Job 與回報狀態。'),
        ('通知 Sending 長時間未改變','不能直接當作未寄出；SMTP 接受後回寫中斷可能留下 Sending。查排程、資料庫及郵件系統紀錄，避免重複寄送。'),
        ('只有 Started 或沒有新 Log','可能沒有待處理 Job，也可能 Worker／依賴異常；查既有工作與時間範圍，不能只憑 Log 稀少判定當機。')
    ],[2.1,4.9])
    p('收集證據時遮蔽帳密、Cookie、信箱、姓名、帳號、內部路徑及業務識別；保留可在授權環境對照的代號與時間順序。完整診斷資料格式見 [[15.1]]。',kind='維運建議')


def overview(path):
    im=Image.new('RGB',(1600,750),'white');d=ImageDraw.Draw(im)
    title=ImageFont.truetype('C:/Windows/Fonts/msjh.ttc',40)
    font=ImageFont.truetype('C:/Windows/Fonts/msjh.ttc',29)
    d.text((40,20),'三個服務的責任與依賴',font=title,fill='#1f4e78')
    entries=[('ICP 網站','使用者操作與卡控\n保存資料與建立請求\n報表、附件、通知','ICP DB／ILC／FIESTA\nRabbitMQ／Hub HTTP\nSMTP／檔案目錄'),
             ('TEL.IntegrationHub','接收起案事件\n協調 Export Job\n記錄處理及完成回寫','RabbitMQ／Hub DB\nICP DB／ILC 目的資料\nHTTP 回報入口'),
             ('ICPFileGenerator','背景輪詢與 Claim Job\n讀取輸入及產生 Excel/PDF\n回報 Hub 處理結果','ICP DB 的 Job／查表\n輸入與輸出檔案目錄\nHub HTTP')]
    for i,(name,role,dep) in enumerate(entries):
        x=35+i*525
        d.rounded_rectangle((x,100,x+495,370),radius=14,fill='#eaf1f8',outline='#1f4e78',width=3)
        d.text((x+18,120),name,font=title,fill='#1f4e78')
        d.multiline_text((x+18,190),role,font=font,fill='#172d45',spacing=14)
        d.line((x+247,372,x+247,425),fill='#1f4e78',width=4)
        d.text((x+155,384),'使用的依賴',font=font,fill='#172d45')
        d.rounded_rectangle((x,440,x+495,650),radius=14,fill='#f4f5f7',outline='#778899',width=2)
        d.multiline_text((x+18,460),dep,font=font,fill='#172d45',spacing=14)
    d.text((40,690),'圖例：上下連線表示依賴，不表示順序、實體主機或資料庫外鍵。',font=font,fill='#172d45')
    im.save(path)


def architecture(doc,table,p,h,root,flow_image):
    assets=root/'outputs/manual/admin-v24-assets';assets.mkdir(parents=True,exist_ok=True)
    doc.add_heading('19 架構流程圖',level=1).paragraph_format.page_break_before=True
    p('先看三個服務各做什麼，再看起案、產檔及通知三條流程。圖是目前程式責任與資料流的示意，不代表主機部署拓樸或已完成外部系統實測。使用者已告知 Hub 與 FileGenerator 已啟動，本章不要求重新啟動。',kind='目前程式行為')
    overview(assets/'overview.png')
    doc.add_paragraph().add_run().add_picture(str(assets/'overview.png'),width=Inches(6.7))
    table(['容易混淆的地方','如何理解'],[
        ('三服務與 IIS','ICP 提供使用者畫面；Hub 處理整合；FileGenerator 做背景產檔。IIS 是承載方式，不等於業務狀態；API 可回應不表示工作鏈完成。'),
        ('Job、Outbox、Log','Outbox 保存待發布事件；Job 保存產檔任務；Log 記錄處理或寄送。它們用途不同，不能互相替代完成證據。'),
        ('資料庫與檔案','SQL 保存資料、識別與狀態；Excel／PDF／原檔保存在目錄。資料列存在不代表檔案一定存在；路徑需按服務身分與 ContentRoot 核對。')
    ],[1.8,5.2])
    specs=[
      ('19.2 起案流程圖與說明','case.png','起案 由 ICP 到 Hub 再到目的資料',[
       '1 ICP 驗證起案\nHeader／Detail 快照','2 保存 Outbox\n同一筆事件識別','3 Outbox Worker\n發布 RabbitMQ','4 Hub Consumer\n依 Deposit／ARUR','5 寫目的資料\n記錄 MESSAGE_LOG','6 回寫 ICP\n事件／案件狀態'],[
       ('1 驗證與快照','Controller／Service 檢查必填及起案條件，把當下 Header、Detail 等內容組成事件；預覽不是送出成功。'),
       ('2–3 非同步交接','Outbox 保存後由背景 Worker 發布；RabbitMQ 是事件傳遞，Published 只表示發布階段完成。停用發布時事件可仍是 Pending。'),
       ('4–5 寫入與紀錄','Hub Consumer 辨別 Deposit／ARUR，呼叫專用 ILC 寫入服務並記錄成功或失敗；不是由 FileGenerator 處理起案。'),
       ('6 完成判斷','以同一 messageId 串起 Outbox、Hub Log 與下游結果，確認完成回寫；不要只看 Hub 程序已啟動。')],
       '圖中箭頭代表事件處理方向；下游部分寫入或回寫失敗的恢復界線見 [[21.1]]、[[24.1]]、[[26.4]]。來源：ICP Services/Integration、Hub Consumers 與 IcpOutboxCompletionService。'),
      ('19.3 Export 產檔流程圖與說明','export.png','Export 由請求到檔案下載',[
       '1 ICP 保存輸入\n建立 Export Request','2 HTTP 通知 Hub\n傳送 RequestId','3 Hub 建立 Job\n同環境 ICP DB','4 FileGenerator\nClaim Job 讀檔查表','5 產生輸出檔\n保存 Job 結果','6 HTTP 回報 Hub\nRequest 回寫與下載'],[
       ('1–2 使用者入口','ICP 保存 Shipping advice 輸入檔與 Export Request，透過 HTTP 交給 Hub；目前這條主流程不是靠起案 RabbitMQ Outbox。'),
       ('3–4 背景執行','Hub 建立 Job；FileGenerator 從同環境資料庫輪詢並 Claim。RequestId 是請求識別，JobId 是工作識別，兩者透過 Job.RequestId 關聯。'),
       ('5 輸出位置','Reader 與查表服務準備資料，Generator 產生 Excel／PDF 至輸出目錄；不同服務需能讀到同一實體位置。'),
       ('6 完成判斷','Job 完成後 HTTP 回報 Hub，Hub 更新 Request 供 ICP 查詢下載。核對 Job、Request、實際檔案及 callback；有檔案但回寫失敗時不要立即重建。')],
       '箭頭是請求、資料庫工作與回報的順序，不表示 Hub 與 FileGenerator 只有單向 HTTP。排查及狀態見 [[11]]、[[21.2]]；Log 串聯見 [[15.3]]。'),
      ('19.4 通知信流程圖與說明','mail.png','通知 由觸發到寄送紀錄',[
       '1 條件／使用者觸發\n依通知種類處理','2 Worker 或通知服務\n讀設定與排程','3 先保存 Log\nState = Sending','4 共用 SMTP 元件\n送到郵件伺服器','5 回寫寄送結果\nSent／Failed','6 維運查核\n信箱投遞另行確認'],[
       ('1–2 觸發與設定','到貨延遲、Arrival Notice、管制品與押金通知各有觸發邏輯；並非全部都排同一張排程表。依 Enabled、時段、條件及收件者來源決定。'),
       ('3 先記錄再寄送','NotificationMailSender 先保存 NOTIFICATION_MAIL_LOG 的 Sending；無法保存本次嘗試時不先寄出郵件。'),
       ('4–5 SMTP 與結果','共用 NotificationSmtpClient 建立及送出信件；一般寄送例外寫 Failed，正常回寫 Sent／SentUtc。取消或回寫失敗仍需分段核對。'),
       ('6 圖中的人工步驟','最後一格是維運建議，不是程式保證投遞到信箱。Sent 不等於對方已收信；SMTP 接受後回寫中斷也可能造成重送風險。')],
       '此流程在 ICP 通知服務內處理，不經 FileGenerator 或一律透過 Hub。名單修改見 [[12.1]]、排查見 [[21.3]]，時間與避免重複見 [[12.2]]。')]
    for title,name,caption,nodes,rows,note in specs:
        h(title).paragraph_format.page_break_before=True
        flow_image(assets/name,caption,nodes)
        doc.add_paragraph().add_run().add_picture(str(assets/name),width=Inches(6.7))
        table(['圖中步驟','閱讀重點'],rows,[1.65,5.35])
        p(note,kind='維運建議')


def screenshot_pages(doc,table,p,h,root):
    h('26.8 管理畫面與程式對照').paragraph_format.page_break_before=True
    p('本章為使用者重啟系統後，於 2026-10-27 補拍的本機管理畫面，並依 Controller／View／Service 核對用途。只查詢與截圖，未新增、刪除、修改權限、上傳資料或寄信；裁切及遮蔽範圍依各圖說明。',kind='維運建議')
    assets=root/'outputs/manual/admin-v24-assets';assets.mkdir(exist_ok=True)
    specs=[('角色與資源權限','screenshots-admin-20261027/role-permissions.png',(280,62,1700,710),None,'2026-10-27 本機實拍，裁切省略側欄與登入資訊。',[
        '先看 RoleCode／RoleName，再看 ResourceCode、ResourceType 與操作；頁面權限與按鈕權限須分開理解。',
        '圖中是既有列表，不是建議把 System 角色授予所有人。角色指派、DataScope 與 Session 的完整邏輯見 [[9.2]]。',
        '維護入口：RolePermissionsController、Views/RolePermissions、UserResourcePermissionService。畫面可見不替代後端授權。']),
      ('通知設定為唯讀檢視','screenshots-admin-20261027/notification.png',(290,62,1690,720),'mail','2026-10-27 本機實拍；收件者與副本欄已遮蔽。',[
        '先看頁面提示、SMTP 設定及各通知種類。圖上 Enabled／時段只是拍攝當時值，不是所有環境固定值。',
        'MailTo／CcTo 需交由管理者修改有效 appsettings；ArrivalNotice 的 To 另取 Ship Info 欄位。下方分頁是清單與預覽，預覽不會寄信。',
        '維護入口：NotificationMailOptions、通知頁 Controller／View、NotificationMailSender；名單實例見 [[12.1]]。']),
      ('上傳紀錄與原檔狀態','screenshots-admin-20261027/history.png',(280,62,1700,445),'history','2026-10-27 本機實拍；上傳人帳號已遮蔽。',[
        '先分清 Upload Check 與 Upload History；後者列已成功儲存的批次，可依各欄 Filter 查詢。',
        'Imported／已匯入是資料匯入狀態；File not found／檔案不存在是目前原檔狀態。圖中三筆不能據此判成當初匯入失敗。',
        '維護入口：ForwarderHistoryTableMetadataProvider、Config/forwarder-upload-history-table-fields.json、上傳歷史查詢與下載動作。']),
      ('Settings 欄位與 Excel 入口','screenshots-admin-20261027/deliveryto.png',(280,62,1700,410),'delivery','2026-10-27 本機實拍；地址、城市、電話及區域欄已遮蔽。',[
        '右上方下載／上傳 Excel 分別依查看／新增權限；上傳先驗證整份資料，再取代本頁範圍。',
        '列表的 Key、Customer Code、Name 是不同欄位。Add DI/SA 的 DELIVERY_TO 比對 Customer Code；不要改成 Name 或 Key。',
        '維護入口：SystemConfigControllerBase、SettingsExcelHelper 與 DeliveryToList；對照程式比對值，不從名稱猜測。'])]
    for i,(title,source,box,mask,caption,notes) in enumerate(specs):
        if i:h('圖例 '+str(i+1)+' '+title).paragraph_format.page_break_before=True
        else:h('圖例 1 '+title)
        im=Image.open(root/'outputs/manual'/source).convert('RGB');d=ImageDraw.Draw(im)
        if mask=='mail':
            d.rectangle((978,422,1683,585),fill='#e2e6ea')
        elif mask=='history':d.rectangle((925,315,1115,432),fill='#e2e6ea')
        elif mask=='delivery':d.rectangle((992,212,1699,400),fill='#e2e6ea')
        crop=im.crop(box);dest=assets/f'screen-{i+1}.png';crop.save(dest)
        width=6.7;height=width*crop.height/crop.width
        if height>3.6:width*=3.6/height
        doc.add_paragraph().add_run().add_picture(str(dest),width=Inches(width))
        p(caption,kind='維運建議')
        for j,note in enumerate(notes,1):p(str(j)+'. '+note,kind='維運建議')


def doc_update_checklist(table,p,h):
    h('18.4 文件更新檢查表').paragraph_format.page_break_before=True
    p('依改修內容勾選受影響的文件，不要求每次重寫所有章節。獨立附件「ICP文件更新檢查表.docx」可直接填寫；維護紀錄使用「ICP維護紀錄範本.docx」。兩份是本次交付的空白範本，不含未執行的完成紀錄。',kind='維運建議')
    table(['變更觸發','必須同步檢查'],[
        ('UI／按鈕／欄位順序','操作步驟、畫面截圖、快速索引、欄位名稱與操作權限；截圖註日期與去識別化範圍。'),
        ('Excel／必填／格式／比對','範本與實際下載檔、卡控章節、成功／失敗案例、錯誤文字與本地測試。'),
        ('Config／排程／通知名單','有效來源、生效方式、既有快照影響；保留其他 MailTo／CcTo，更新通知管理說明。'),
        ('SQL／資料模型／整合','Schema、部署順序、相容性、資料流與識別、回復界線；不能只更新 Entity。'),
        ('背景服務／產檔','啟動依賴、Worker／Job 狀態、Log 判讀、路徑與 callback；未實作的 health／retry 仍要標示。'),
        ('文件發布','版本、日期、目錄、交叉連結、頁碼、Word／PDF、網站 docs 副本與 CustomLinks；核對 SHA-256。')
    ],[2.0,5.0])
    p('本輪完成條件：五項使用便利性改善與架構圖說明已納入、附件可填寫、內容與來源一致、頁面完整、連結可解析、下載副本相同。不將仍可想像的新功能列作本輪待辦；後續只在程式或需求改變時維護。',kind='維運建議')


def make_templates(out):
    def base(title):
        doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65)
        for name in ['Normal','Title','Heading 1','Heading 2']:
            style=doc.styles[name];style.font.name='Microsoft JhengHei';style.font.color.rgb=RGBColor(0,0,0)
            style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
        doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
        doc.styles['Title'].font.size=Pt(25)
        for border in list(doc.styles['Title']._element.get_or_add_pPr().findall(qn('w:pBdr'))):
            doc.styles['Title']._element.get_or_add_pPr().remove(border)
        footer=sec.footer.paragraphs[0]
        footer.add_run(title+'　版本 1.0　｜　第 ')
        for instr in ['PAGE','NUMPAGES']:
            fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),instr);footer._p.append(fld)
            footer.add_run(' 頁／共 ' if instr=='PAGE' else ' 頁')
        doc.add_paragraph(title,style='Title');doc.add_paragraph('空白可填寫範本　版本 1.0　2026-10-27')
        doc.core_properties.title=title;doc.core_properties.author='ICP Project Team'
        return doc
    def tab(doc,rows):
        t=doc.add_table(rows=0,cols=2);t.autofit=False
        for label,value in rows:
            c=t.add_row().cells;c[0].text=label;c[1].text=value;c[0].width=Inches(1.5);c[1].width=Inches(5.1)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'eaf1f8');c[0]._tc.get_or_add_tcPr().append(shade)
            t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        return t
    doc=base('ICP維護紀錄範本')
    doc.add_paragraph('每次改修另存一份。以本地端測試作為完成依據；未執行項目填原因，不填成已完成。不貼帳密、完整連線字串、實際收件名單或個資。')
    doc.add_heading('一 需求與修改範圍',level=1)
    tab(doc,[(x,'待填：____________________________') for x in ['需求／BUG 編號','日期與時區／負責角色','問題與修改後行為','服務與版本／commit','Config／SQL／範本影響']])
    doc.add_heading('二 交付檔案',level=1)
    doc.add_paragraph('逐檔填寫：服務／資料夾／檔名／用途／部署順序。原始碼、發布產物、環境設定、SQL、範本與文件分開列。')
    tab(doc,[('資料夾與檔案 1','待填：____________________________'),('資料夾與檔案 2','待填：____________________________'),('其他清單附件','待填：____________________________')])
    doc.add_heading('三 本地測試與文件核對',level=1).paragraph_format.page_break_before=True
    tab(doc,[(x,'待填：____________________________') for x in ['測試專案與指令／組態','輸入及暫存輸出位置','Exit Code／項數／Log','首次失敗與修正結果','文件版本與修改章節','頁碼／連結／副本雜湊']])
    doc.add_heading('四 回復與完成範圍',level=1)
    tab(doc,[(x,'待填：____________________________') for x in ['前版程式／設定位置','Schema 相容及資料保留','Job／事件／通知重複風險','已完成／不適用／未執行','部署是否執行及原因']])
    doc.add_paragraph('檢核：□ 本地結果有證據　□ 未執行事項如實列明　□ 其他 MailTo／CcTo 保留　□ 文件與下載來源同步')
    doc.save(out/'ICP維護紀錄範本.docx')
    doc=base('ICP文件更新檢查表')
    tab(doc,[('需求編號／日期','待填：____________________________'),('文件版本／核對角色','待填：____________________________')])
    doc.add_paragraph('每項填「已完成／不適用及原因」。只核對本次改修範圍；本表不要求正式環境驗收。')
    doc.add_heading('一 內容與來源',level=1)
    for text in ['UI 名稱、按鈕與欄位順序對應程式。','卡控條件、範本、錯誤訊息與成功／失敗例子一致。','通知名單來源、時段與日期欄位說明正確。','Config 生效方式、既有快照影響與權限明列。','SQL、資料模型、流程圖與回復界線同步。','目前程式行為、維運建議、尚未實作明確區分。']:
        doc.add_paragraph('□ '+text+'　結果／原因：________')
    doc.add_heading('二 圖片與本地證據',level=1)
    for text in ['截圖註明日期、用途、裁切與去識別化範圍。','每張架構圖說明元件、箭頭、資料與完成判斷。','Log 範例明示教學示例，不當作實際運行紀錄。','測試專案、組態、指令、Exit Code 與 Log 可追溯。']:
        doc.add_paragraph('□ '+text+'　結果／原因：________')
    doc.add_heading('三 文件發布與收尾',level=1).paragraph_format.page_break_before=True
    for text in ['版本、日期、目錄與交叉連結更新，沒有空白占位連結。','Word／PDF 逐頁核對，表格與圖片沒有裁切或孤立續頁。','網站 docs 與交付原檔雜湊相同，連結版本參數更新。','環境設定只改本次區段，其他名單與設定保留。','部署清單依資料夾與檔案列出；測試工具不整批放網站。','維護紀錄與回復說明完成，未執行事項明列。']:
        doc.add_paragraph('□ '+text+'　結果／原因：________')
    tab(doc,[('核對紀錄或附件路徑','待填：____________________________'),('剩餘問題與處理結果','待填：無／________________________'),('本輪結案範圍','待填：____________________________')])
    doc.add_paragraph('完成原則：本次需求與檢查項已處理即可結案；不以不斷新增同類建議延長文件修改。新功能或新需求另行記錄。')
    doc.save(out/'ICP文件更新檢查表.docx')
